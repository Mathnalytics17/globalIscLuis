import json
from datetime import date

from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.misc.api.models.companies.index import Empresa
from apps.misc.api.models.dynamicTechnicalConfig.index import (
    CatalogoTecnico,
    CatalogoTecnicoItem,
    PruebaFuenteLimite,
    PruebaLimiteCampo,
    CriterioEvaluacionLimite,
)
from apps.misc.api.models.pruebas.index import (
    Prueba,
    PruebaResultado,
    PruebaResultadoComponente,
    PruebaResultadoDivision,
)
from apps.muestras.api.models.muestras.index import Muestra
from apps.muestras.api.models.loteMuestras.index import LoteMuestras
from apps.muestras.api.models.muestraAtributoTecnico.index import MuestraAtributoTecnico
from apps.muestras.api.models.pruebasMuestra.index import PruebaMuestra
from apps.resultado.api.models.index import HistoricoResultado, Resultado, VersionInterpretacion
from apps.muestras.api.models.auditoria import EventoAuditoriaOperativa
from apps.resultado.api.views.index import _build_predefined_comments, _sample_review_status
from apps.users.api.models.index import User


class ResultEntryTests(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="Global Oil Resultado Test")
        self.user = User.objects.create_user(
            email="resultado-entry@example.com",
            password="test-password",
            empresa=self.empresa,
            is_active=True,
            is_superuser=True,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

        self.lote = LoteMuestras.objects.create(
            tipo_cliente="registrado",
            cliente_empresa=self.empresa,
            fecha_envio=date.today(),
            fecha_recepcion=date.today(),
            estado="en_laboratorio",
            usuario_registro=self.user,
        )

        self.muestra = Muestra.objects.create(
            lote=self.lote,
            fecha_toma=timezone.now(),
            usuario_registro=self.user,
            tipo_muestra="aceite",
            is_ingresado=True,
        )
        self.prueba = Prueba.objects.create(
            nombre_variable="Espuma",
            acronimo="FOAM",
        )
        self.sample_test = PruebaMuestra.objects.create(
            muestra=self.muestra,
            prueba=self.prueba,
            usuario_solicitud=self.user,
        )

    def test_sample_status_uses_worst_technical_severity(self):
        tests = [
            {"estado_tecnico": "FUERA_DE_LIMITE", "is_revisada": False, "resultado": {}},
            {"estado_tecnico": "CRITICO", "is_revisada": False, "resultado": {}},
        ]

        self.assertEqual(_sample_review_status(self.muestra, tests), "Crítica")

    def test_predefined_comments_only_come_from_technical_field_configuration(self):
        source = PruebaFuenteLimite.objects.create(
            prueba=self.prueba,
            tipo_limite="global",
        )
        normal_field = PruebaLimiteCampo.objects.create(
            fuente_limite=source,
            nombre="Valor normal",
            codigo="valor_normal",
            comentario_normal="Texto normal configurado.",
        )
        warning_field = PruebaLimiteCampo.objects.create(
            fuente_limite=source,
            nombre="Valor fuera",
            codigo="valor_fuera",
            comentario_fuera_limite="Texto de seguimiento configurado.",
        )
        no_text_field = PruebaLimiteCampo.objects.create(
            fuente_limite=source,
            nombre="Sin texto",
            codigo="sin_texto",
        )

        comments = _build_predefined_comments([{
            "evaluacion_limite": {
                "details": [
                    {"field_id": normal_field.id, "estado": "NORMAL"},
                    {"field_id": warning_field.id, "estado": "CRITICO"},
                    {"field_id": no_text_field.id, "estado": "NORMAL"},
                ],
            },
        }])

        self.assertEqual(comments, [
            "Texto normal configurado.",
            "Texto de seguimiento configurado.",
        ])

    def test_invalidated_result_has_priority_over_stale_review_flags(self):
        tests = [
            {
                "estado_tecnico": "CRITICO",
                "is_revisada": True,
                "resultado": {"id": 999, "estatus": "rechazado"},
            },
        ]

        self.assertEqual(_sample_review_status(self.muestra, tests), "Requiere corrección")

    def test_invalidating_result_unreviews_and_blocks_the_sample(self):
        self.sample_test.completada = True
        self.sample_test.is_revisada = True
        self.sample_test.estatus = "aprobado"
        self.sample_test.save(update_fields=["completada", "is_revisada", "estatus"])
        self.muestra.is_resultado_ingresado = True
        self.muestra.is_revisado = True
        self.muestra.save(update_fields=["is_resultado_ingresado", "is_revisado"])
        result = Resultado.objects.create(
            prueba_muestra=self.sample_test,
            resultado=json.dumps({"values": [{"key": "valor", "value": "12"}]}),
            fecha_medicion=timezone.now(),
            usuario_medicion=self.user,
            estatus="aprobado",
        )

        response = self.client.post(
            f"/api/lubrication/results/{result.id}/invalidate/",
            {"reason": "Control de calidad."},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.sample_test.refresh_from_db()
        self.muestra.refresh_from_db()
        self.assertEqual(self.sample_test.estatus, "rechazado")
        self.assertFalse(self.sample_test.completada)
        self.assertFalse(self.sample_test.is_revisada)
        self.assertFalse(self.muestra.is_resultado_ingresado)
        self.assertFalse(self.muestra.is_revisado)
        self.assertEqual(self.muestra.campos_adicionales["interpretacion"]["estado"], "requiere_revision")

    def test_consolidated_export_accepts_timezone_aware_sample_date(self):
        response = self.client.get(
            "/api/lubrication/result-entry/excel/export/",
            {"lote": self.lote.id, "modo": "consolidado"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_sample_columns_export_includes_value_status_and_rule_columns(self):
        response = self.client.get(
            "/api/lubrication/result-entry/excel/export/",
            {"lote": self.lote.id, "modo": "por_muestra_columnas", "muestra": self.muestra.id},
        )

        self.assertEqual(response.status_code, 200)
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        headers = [cell.value for cell in workbook.active[1]]
        self.assertTrue(any(str(value).startswith("Estado |") for value in headers))
        self.assertTrue(any(str(value).startswith("Regla |") for value in headers))

    def test_upload_template_supports_both_layouts_and_typed_validation(self):
        result = PruebaResultado.objects.create(prueba=self.prueba, nombre="Presencia")
        division = PruebaResultadoDivision.objects.create(resultado=result, es_principal=True)
        PruebaResultadoComponente.objects.create(
            division=division,
            nombre="Presencia",
            acronimo="P",
            tipo_dato="booleano",
            etiqueta_verdadero="Hay",
            etiqueta_falso="No hay",
        )

        from io import BytesIO
        from openpyxl import load_workbook

        consolidated = self.client.get(
            "/api/lubrication/result-entry/excel/template/",
            {"lote": self.lote.id, "modo": "consolidado"},
        )
        self.assertEqual(consolidated.status_code, 200, getattr(consolidated, "data", None))
        workbook = load_workbook(BytesIO(consolidated.content))
        self.assertIn("Resultados", workbook.sheetnames)
        self.assertEqual(workbook.active.title, "Resultados")
        sheet = workbook["Resultados"]
        self.assertTrue(sheet.protection.sheet)
        self.assertTrue(sheet["A2"].protection.locked)
        self.assertFalse(sheet["L2"].protection.locked)
        self.assertTrue(any("L2" in str(validation.sqref) for validation in sheet.data_validations.dataValidation))

        per_sample = self.client.get(
            "/api/lubrication/result-entry/excel/template/",
            {"lote": self.lote.id, "modo": "por_muestra"},
        )
        self.assertEqual(per_sample.status_code, 200, getattr(per_sample, "data", None))
        per_sample_workbook = load_workbook(BytesIO(per_sample.content))
        self.assertIn(str(self.muestra.id), per_sample_workbook.sheetnames)
        self.assertEqual(per_sample_workbook.active.title, str(self.muestra.id))
        # Excel puede entregar un booleano nativo (TRUE/FALSE), no solo texto.
        per_sample_workbook[str(self.muestra.id)]["L2"] = True
        per_sample_output = BytesIO()
        per_sample_workbook.save(per_sample_output)
        from django.core.files.uploadedfile import SimpleUploadedFile
        per_sample_upload = SimpleUploadedFile(
            "resultados_por_muestra.xlsx",
            per_sample_output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        preview = self.client.post(
            "/api/lubrication/result-entry/excel/preview/",
            {"file": per_sample_upload},
            format="multipart",
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.data["valid"])
        self.assertEqual(preview.data["rows_valid"], 1)
        self.assertEqual(
            preview.data["values_by_test"][str(self.sample_test.id)][0]["value"],
            "true",
        )

        # La etiqueta visible configurada debe convertirse al mismo valor canonico
        # para que el selector web pueda mostrarla inmediatamente.
        per_sample_workbook[str(self.muestra.id)]["L2"] = "Hay"
        label_output = BytesIO()
        per_sample_workbook.save(label_output)
        label_upload = SimpleUploadedFile(
            "resultados_boolean_label.xlsx",
            label_output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        label_preview = self.client.post(
            "/api/lubrication/result-entry/excel/preview/",
            {"file": label_upload},
            format="multipart",
        )
        self.assertEqual(label_preview.status_code, 200, label_preview.data)
        self.assertTrue(label_preview.data["valid"])
        self.assertEqual(
            label_preview.data["values_by_test"][str(self.sample_test.id)][0]["value"],
            "true",
        )

    def test_excel_preview_accepts_valid_xlsx_without_worksheet_dimensions(self):
        """XLSX producers may omit the optional worksheet dimension metadata."""
        from io import BytesIO
        import re
        from zipfile import ZIP_DEFLATED, ZipFile

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook

        template = self.client.get(
            "/api/lubrication/result-entry/excel/template/",
            {"lote": self.lote.id, "modo": "por_muestra"},
        )
        self.assertEqual(template.status_code, 200, getattr(template, "data", None))
        workbook = load_workbook(BytesIO(template.content))
        workbook[str(self.muestra.id)]["L2"] = "12.5"
        regular_output = BytesIO()
        workbook.save(regular_output)

        no_dimensions = BytesIO()
        with ZipFile(BytesIO(regular_output.getvalue())) as source, ZipFile(no_dimensions, "w", ZIP_DEFLATED) as destination:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                    content = re.sub(br"<dimension[^>]*/>", b"", content)
                destination.writestr(item, content)

        upload = SimpleUploadedFile(
            "resultados_sin_dimension.xlsx",
            no_dimensions.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        preview = self.client.post(
            "/api/lubrication/result-entry/excel/preview/",
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.data["valid"])
        self.assertEqual(preview.data["rows_valid"], 1)

    def test_excel_import_applies_validated_values_as_draft(self):
        template = self.client.get(
            "/api/lubrication/result-entry/excel/template/",
            {"lote": self.lote.id, "modo": "consolidado"},
        )
        self.assertEqual(template.status_code, 200)

        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(template.content))
        workbook["Resultados"]["L2"] = "12.5"
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            "resultados.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        preview = self.client.post(
            "/api/lubrication/result-entry/excel/preview/",
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.data["valid"])
        self.assertEqual(preview.data["rows_valid"], 1)

        output.seek(0)
        upload = SimpleUploadedFile(
            "resultados.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        imported = self.client.post(
            "/api/lubrication/result-entry/excel/import/",
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(imported.data["rows_imported"], 1)
        self.sample_test.refresh_from_db()
        self.assertEqual(self.sample_test.estatus, "en_proceso")
        result = Resultado.objects.get(prueba_muestra=self.sample_test)
        self.assertEqual(json.loads(result.resultado)["values"][0]["value"], "12.5")

    def test_confirm_all_batch_results_promotes_complete_drafts(self):
        form = self.client.get(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/form/"
        )
        self.assertEqual(form.status_code, 200, form.data)
        field_key = form.data["fields"][0]["key"]
        draft = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/save-draft/",
            {"values": [{"key": field_key, "value": "12.5", "observacion": ""}]},
            format="json",
        )
        self.assertEqual(draft.status_code, 200, draft.data)

        response = self.client.post(
            f"/api/lubrication/result-entry/batches/{self.lote.id}/confirm-all/",
            {},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["tests_completed"], 1)
        self.sample_test.refresh_from_db()
        self.assertEqual(self.sample_test.estatus, "completado")
        self.assertTrue(self.sample_test.completada)
        result = Resultado.objects.get(prueba_muestra=self.sample_test)
        self.assertEqual(result.estatus, "preliminar")

    def test_repeating_the_same_correction_confirmation_is_idempotent(self):
        reviewed_test = Prueba.objects.create(
            nombre_variable="Ensayo ya revisado",
            acronimo="REVIEWED",
        )
        PruebaMuestra.objects.create(
            muestra=self.muestra,
            prueba=reviewed_test,
            usuario_solicitud=self.user,
            completada=True,
            is_revisada=True,
            estatus="aprobado",
        )
        Resultado.objects.create(
            prueba_muestra=self.sample_test,
            resultado=json.dumps({"values": []}),
            fecha_medicion=timezone.now(),
            usuario_medicion=self.user,
            estatus="rechazado",
        )
        self.sample_test.estatus = "rechazado"
        self.sample_test.save(update_fields=["estatus"])
        form = self.client.get(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/form/"
        )
        field_key = form.data["fields"][0]["key"]
        payload = {
            "idempotency_key": "correccion-repetida-1",
            "values": [{"key": field_key, "value": "12.5", "observacion": ""}],
        }

        first = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/confirm/",
            payload,
            format="json",
        )
        repeated = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/confirm/",
            payload,
            format="json",
        )

        self.assertEqual(first.status_code, 200, first.data)
        self.assertEqual(repeated.status_code, 200, repeated.data)
        self.assertTrue(repeated.data["idempotent"])

    def test_interpretation_trends_orders_by_existing_test_fields(self):
        form = self.client.get(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/form/"
        )
        field_key = form.data["fields"][0]["key"]
        confirmed = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/confirm/",
            {"values": [{"key": field_key, "value": "12.5", "observacion": ""}]},
            format="json",
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.data)

        response = self.client.get(
            f"/api/lubrication/interpretation/samples/{self.muestra.id}/trends/",
            {"pruebas": self.prueba.acronimo},
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn("groups", response.data)

    def test_saving_interpretation_creates_an_immutable_version_and_event(self):
        # La interpretación solo se guarda después de revisar todos los ensayos.
        self.muestra.is_revisado = True
        self.muestra.save(update_fields=["is_revisado"])
        response = self.client.patch(
            f"/api/lubrication/interpretation/samples/{self.muestra.id}/",
            {
                "conclusion": "El activo puede continuar en operación.",
                "comentarios_predefinidos": ["Sin condición crítica."],
                "comentarios_manuales": [],
                "graficas_activas": True,
                "graficas_seleccionadas": [self.prueba.acronimo],
                "motivo": "Primera interpretación técnica.",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        version = VersionInterpretacion.objects.get(muestra=self.muestra)
        self.assertEqual(version.version, 1)
        self.assertEqual(version.contenido["conclusion"], "El activo puede continuar en operación.")
        self.assertTrue(EventoAuditoriaOperativa.objects.filter(
            accion="interpretacion.guardada",
            muestra=self.muestra,
            entidad_id=str(version.id),
        ).exists())

    def test_result_history_reconstructs_each_correction_transition(self):
        def payload(value):
            return json.dumps({"values": [{"key": "valor", "value": value, "value_label": value}]})

        result = Resultado.objects.create(
            prueba_muestra=self.sample_test,
            resultado=payload("3"),
            fecha_medicion=timezone.now(),
            usuario_medicion=self.user,
            estatus="preliminar",
        )
        for previous in ("1", "2"):
            HistoricoResultado.objects.create(
                resultado=result,
                resultado_anterior=payload(previous),
                fecha_medicion_anterior=result.fecha_medicion,
                usuario_medicion=self.user,
                usuario_modificacion=self.user,
                motivo_cambio=f"Corrección hacia el valor siguiente desde {previous}.",
            )

        response = self.client.get(
            f"/api/lubrication/review-results/results/{result.id}/history/"
        )

        self.assertEqual(response.status_code, 200, response.data)
        corrections = [item for item in response.data["history"] if item["type"] == "correction"]
        self.assertEqual([(item["valor_anterior"], item["valor_nuevo"]) for item in corrections], [("1", "2"), ("2", "3")])

    def test_confirm_composite_result_keeps_full_payload_and_short_sample_value(self):
        values = []
        for sequence, component_names in [
            ("I", ["FI", "EI"]),
            ("II", ["FII", "EII"]),
            ("III", ["FIII", "EIII"]),
        ]:
            result = PruebaResultado.objects.create(prueba=self.prueba, nombre=f"Secuencia {sequence}")
            division = PruebaResultadoDivision.objects.create(resultado=result, es_principal=True)
            for component_name in component_names:
                component = PruebaResultadoComponente.objects.create(
                    division=division,
                    nombre=component_name,
                    acronimo=component_name,
                )
                values.append({
                    "key": f"resultado:{result.id}:division:{division.id}:componente:{component.id}",
                    "value": str(10 + len(values)),
                    "observacion": "",
                })

        response = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/confirm/",
            {
                "values": values,
                "observaciones": "Resultado compuesto de espuma",
                "fecha_medicion": "2026-06-03T22:22:42.917Z",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.sample_test.refresh_from_db()
        self.assertEqual(self.sample_test.valor, "10 / 11 / 12 / 13 / 14 / 15")
        self.assertLessEqual(len(self.sample_test.valor), 255)

        result = Resultado.objects.get(prueba_muestra=self.sample_test)
        payload = json.loads(result.resultado)
        self.assertEqual(len(payload["values"]), 6)
        self.assertEqual(payload["values"][0]["key"], values[0]["key"])

    def test_confirm_composite_result_persists_component_limit_evaluation(self):
        catalog = CatalogoTecnico.objects.create(
            nombre="Nivel de desempeno",
            codigo="nivel_de_desempeno",
            tipo_muestra="aceite",
        )
        catalog_item = CatalogoTecnicoItem.objects.create(catalogo=catalog, nombre="API SL")
        MuestraAtributoTecnico.objects.create(
            muestra=self.muestra,
            catalogo=catalog,
            item=catalog_item,
        )
        source = PruebaFuenteLimite.objects.create(
            prueba=self.prueba,
            catalogo_fuente=catalog,
            politica_evaluacion="todas_deben_cumplir",
        )

        values = []
        criterion_values = {}
        for sequence, component_names, limits in [
            ("1", ["EI", "FI"], {"EI": "11", "FI": "0"}),
            ("2", ["EII", "FII"], {"EII": "25", "FII": "0"}),
            ("3", ["EIII", "FIII"], {"EIII": "11", "FIII": "0"}),
        ]:
            result = PruebaResultado.objects.create(prueba=self.prueba, nombre=f"Secuencia {sequence}")
            division = PruebaResultadoDivision.objects.create(resultado=result, es_principal=True)
            for component_name in component_names:
                component = PruebaResultadoComponente.objects.create(
                    division=division,
                    nombre=component_name,
                    acronimo=component_name,
                )
                field = PruebaLimiteCampo.objects.create(
                    fuente_limite=source,
                    resultado=result,
                    nombre=f"Secuencia {sequence} {component_name}",
                    codigo=f"secuencia_{sequence}_{component_name.lower()}",
                    operador="max",
                )
                criterion_values[field.codigo] = limits[component_name]
                values.append({
                    "key": f"resultado:{result.id}:division:{division.id}:componente:{component.id}",
                    "value": "10" if component_name == "FI" else "0",
                    "observacion": "",
                })
        CriterioEvaluacionLimite.objects.create(
            fuente_limite=source,
            nombre=catalog_item.nombre,
            codigo=catalog_item.codigo,
            tipo_criterio="catalogo_item",
            catalogo_item=catalog_item,
            valores_limite=criterion_values,
        )

        response = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{self.sample_test.id}/confirm/",
            {"values": values, "observaciones": "", "fecha_medicion": timezone.now().isoformat()},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.sample_test.refresh_from_db()
        self.assertEqual(self.sample_test.estado_limite, "CRÍTICO")
        self.assertTrue(self.sample_test.evaluacion_limite["details"])
        failing = [
            item for item in self.sample_test.evaluacion_limite["details"]
            if item["field"].endswith("_fi")
        ][0]
        self.assertEqual(failing["limit_value"], "0")
        self.assertFalse(failing["passed"])

    def test_assigned_limit_criterion_evaluates_contamination_without_sample_attribute(self):
        prueba = Prueba.objects.create(
            nombre_variable="Nivel de contaminacion",
            acronimo="NCONT",
        )
        sample_test = PruebaMuestra.objects.create(
            muestra=self.muestra,
            prueba=prueba,
            usuario_solicitud=self.user,
        )
        catalog = CatalogoTecnico.objects.create(
            nombre="Nivel de contaminacion objetivo",
            codigo="nivel_de_contaminacion",
            tipo_muestra="aceite",
        )
        target = CatalogoTecnicoItem.objects.create(catalogo=catalog, nombre="16/14/11")
        sample_test.criterio_limite_catalogo = catalog
        sample_test.criterio_limite_item = target
        sample_test.save(update_fields=["criterio_limite_catalogo", "criterio_limite_item"])

        source = PruebaFuenteLimite.objects.create(
            prueba=prueba,
            catalogo_fuente=catalog,
            politica_evaluacion="todas_deben_cumplir",
        )
        result = PruebaResultado.objects.create(prueba=prueba, nombre="Conteo")
        division = PruebaResultadoDivision.objects.create(resultado=result, es_principal=True)
        values = []
        criterion_values = {}
        for acronym, limit, measured in [("4um", "16", "55"), ("6um", "14", "22"), ("14um", "11", "33")]:
            component = PruebaResultadoComponente.objects.create(
                division=division,
                nombre=acronym,
                acronimo=acronym,
            )
            field = PruebaLimiteCampo.objects.create(
                fuente_limite=source,
                resultado=result,
                componente=component,
                nombre=f"{acronym} maximo",
                codigo=f"nvlcontaminacion_{acronym}_valor_maximo",
                operador="max",
            )
            criterion_values[field.codigo] = limit
            values.append({
                "key": f"resultado:{result.id}:division:{division.id}:componente:{component.id}",
                "value": measured,
                "observacion": "",
            })
        CriterioEvaluacionLimite.objects.create(
            fuente_limite=source,
            nombre=target.nombre,
            codigo=target.codigo,
            tipo_criterio="catalogo_item",
            catalogo_item=target,
            valores_limite=criterion_values,
        )

        response = self.client.post(
            f"/api/lubrication/result-entry/sample-tests/{sample_test.id}/confirm/",
            {"values": values, "observaciones": "", "fecha_medicion": timezone.now().isoformat()},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        sample_test.refresh_from_db()
        self.assertEqual(sample_test.estado_limite, "CRÍTICO")
        self.assertEqual(sample_test.evaluacion_limite["details"][0]["item_label"], "16/14/11")
        self.assertEqual(sample_test.evaluacion_limite["details"][0]["limit_value"], "16")
