import json
import re
import uuid
from collections import defaultdict
from io import BytesIO
from decimal import Decimal, InvalidOperation
from types import SimpleNamespace

from django.http import HttpResponse
from django.db import transaction
from django.db.models import Count, Q
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from apps.utils.pagination import StandardResultsSetPagination
from permissions import ActionPermissionMixin, DenyReadOnlyWrite, HasSecurityPermission
from apps.users.api.models.index import User

from apps.muestras.api.models.loteMuestras.index import LoteMuestras
from apps.muestras.api.models.muestras.index import Muestra
from apps.muestras.api.models.pruebasMuestra.index import PruebaMuestra
from apps.muestras.api.services.limit_engine import (
    resolve_and_evaluate_result,
    resolve_and_evaluate_composite_result,
    resolve_and_evaluate_values_for_sample_test,
    _evaluate_typed,
    _global_limit_reason,
    _global_limit_value_for_field,
    _criterion_item_for_source,
)
from apps.muestras.api.services.workflow import require_batch_operation
from apps.misc.api.models.dynamicTechnicalConfig.index import (
    PruebaFuenteLimite,
    PruebaLimiteCampo,
)
from apps.misc.api.models.technicalCatalogs.index import EquipoPrueba, MetodoEquipo
from apps.muestras.api.serializers.pruebasMuestra.index import PruebaMuestraSerializer
from apps.resultado.api.models.index import Resultado, HistoricoResultado, RevisionResultado
from apps.misc.api.models.pruebas.index import (
    PruebaResultado,
    PruebaResultadoDivision,
    PruebaResultadoComponente,
)
from apps.resultado.api.serializers.index import (
    ResultadoSerializer,
    CreateResultadoSerializer,
    HistoricoResultadoSerializer,
    RevisionResultadoSerializer,
)


def _user_payload(user):
    if not user:
        return None
    full_name = f"{getattr(user, 'first_name', '')} {getattr(user, 'last_name', '')}".strip()
    return {
        "id": getattr(user, "id", None),
        "username": getattr(user, "username", None),
        "email": getattr(user, "email", None),
        "full_name": full_name or getattr(user, "username", None) or getattr(user, "email", None),
    }

def _safe_json_loads(value):
    if value in [None, ""]:
        return None
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:
        return value


def _active_result_for(prueba_muestra):
    prefetched = getattr(prueba_muestra, "_prefetched_objects_cache", {}).get(
        "resultados"
    )
    if prefetched is not None:
        return max(
            prefetched,
            key=lambda item: (
                (
                    item.fecha_actualizacion
                    or item.fecha_registro
                    or item.fecha_medicion
                ).timestamp(),
                item.id,
            ),
            default=None,
        )
    return (
        Resultado.objects
        .filter(prueba_muestra=prueba_muestra)
        .order_by("-fecha_actualizacion", "-id")
        .first()
    )


def _sync_sample_completion_from_tests(muestra):
    if not muestra:
        return
    totals = PruebaMuestra.objects.filter(
        muestra=muestra, estado_asignacion="confirmada"
    ).aggregate(
        total=Count("pk"),
        completed=Count("pk", filter=Q(completada=True)),
        reviewed=Count("pk", filter=Q(is_revisada=True)),
    )
    total = totals["total"]
    completed = totals["completed"]
    reviewed = totals["reviewed"]
    next_result = bool(total and completed == total)
    next_review = bool(total and reviewed == total)
    updates = []
    if muestra.is_resultado_ingresado != next_result:
        muestra.is_resultado_ingresado = next_result
        updates.append("is_resultado_ingresado")
    if muestra.is_revisado != next_review:
        muestra.is_revisado = next_review
        updates.append("is_revisado")
    if updates:
        muestra.save(update_fields=updates)


def _display_value_for_field(field, value):
    if value in [None, ""]:
        return value
    text = str(value)
    normalized = text.strip().casefold()
    for option in field.get("opciones") or []:
        option_value = str(option.get("value", "")).strip().casefold()
        option_label = str(option.get("label", "")).strip()
        if normalized == option_value or normalized == option_label.casefold():
            return option_label or text
    return text


def _compact_result_value(values, max_length=255):
    """
    PruebaMuestra.valor es solo un resumen corto para listados.
    El payload completo vive en Resultado.resultado.
    """
    if not values:
        return None

    parts = []
    for item in values:
        value = item.get("value") if isinstance(item, dict) else item
        if isinstance(item, dict):
            value = item.get("value_label") or value
        if value not in [None, ""]:
            parts.append(str(value))

    if not parts:
        return None

    summary = " / ".join(parts)
    if len(summary) <= max_length:
        return summary

    return summary[: max_length - 3].rstrip() + "..."


def _build_result_fields(prueba_muestra):
    """
    Construye una lista plana de campos a ingresar desde la estructura dinámica de la prueba:
    prueba -> resultados -> divisiones -> componentes.

    Si la prueba no trae componentes, genera un campo simple.
    """
    prueba = prueba_muestra.prueba
    fields = []
    limit_sources = list(
        PruebaFuenteLimite.objects
        .filter(prueba=prueba, activo=True, deleted_at__isnull=True)
        .order_by("prioridad", "id")
    )

    def field_limit(resultado=None, division=None, componente=None):
        for source in limit_sources:
            matches = _matching_limit_fields(source, resultado=resultado, division=division, componente=componente)
            if matches:
                return matches[0]
        return None

    def limit_options(limit_field):
        if not limit_field:
            return []
        if limit_field.tipo_comparacion in ["escala", "escala_ordinal"] and limit_field.escala_comparacion_id:
            return [
                {"value": item.etiqueta, "label": item.etiqueta, "orden": item.orden}
                for item in limit_field.escala_comparacion.items
                .filter(activo=True, deleted_at__isnull=True)
                .order_by("orden", "id")
            ]
        if limit_field.tipo_comparacion == "booleano":
            component = getattr(limit_field, "componente", None)
            return [
                {"value": "true", "label": getattr(component, "etiqueta_verdadero", None) or "Sí"},
                {"value": "false", "label": getattr(component, "etiqueta_falso", None) or "No"},
            ]
        if limit_field.tipo_comparacion in ["opcion", "texto"]:
            options = limit_field.opciones_permitidas
            if isinstance(options, list) and options:
                return [{"value": str(item), "label": str(item)} for item in options]
            evaluation = limit_field.evaluacion_opciones
            if isinstance(evaluation, dict) and evaluation:
                return [{"value": str(item), "label": str(item)} for item in evaluation.keys()]
        return []

    def enrich_with_limit(field_payload, resultado=None, division=None, componente=None):
        limit_field = field_limit(resultado=resultado, division=division, componente=componente)
        if not limit_field:
            return field_payload

        field_payload = {
            **field_payload,
            "limite_campo": limit_field.id,
            "tipo_comparacion": limit_field.tipo_comparacion,
            "operador_limite": limit_field.operador,
            "escala_comparacion": limit_field.escala_comparacion_id or field_payload.get("escala_comparacion"),
        }

        options = limit_options(limit_field)
        if options:
            field_payload["opciones"] = options
            if limit_field.tipo_comparacion in ["escala", "escala_ordinal", "booleano"]:
                field_payload["tipo_dato"] = "escala" if limit_field.tipo_comparacion == "escala_ordinal" else limit_field.tipo_comparacion

        return field_payload

    def component_options(component):
        if component and component.escala_comparacion_id:
            return [
                {"value": item.etiqueta, "label": item.etiqueta, "orden": item.orden}
                for item in component.escala_comparacion.items.filter(activo=True, deleted_at__isnull=True).order_by("orden", "id")
            ]
        if component and component.tipo_dato == "booleano":
            return [
                {"value": "true", "label": component.etiqueta_verdadero or "Si"},
                {"value": "false", "label": component.etiqueta_falso or "No"},
            ]
        return []

    resultados = prueba.resultados.filter(activo=True).order_by("orden", "id")

    for resultado in resultados:
        divisiones = resultado.divisiones.filter(activo=True).order_by("orden", "id")

        if not divisiones.exists():
            fields.append(enrich_with_limit({
                "key": f"resultado:{resultado.id}:valor",
                "resultado_id": resultado.id,
                "resultado_nombre": resultado.nombre,
                "division_id": None,
                "division_nombre": None,
                "componente_id": None,
                "componente_nombre": "Valor",
                "label": resultado.nombre or prueba.nombre_variable,
                "unidad": resultado.unidad_medida or "",
                "tipo_dato": "decimal",
            }, resultado=resultado))
            continue

        for division in divisiones:
            componentes = division.componentes.filter(activo=True).order_by("orden", "id")

            if not componentes.exists():
                fields.append(enrich_with_limit({
                    "key": f"resultado:{resultado.id}:division:{division.id}:valor",
                    "resultado_id": resultado.id,
                    "resultado_nombre": resultado.nombre,
                    "division_id": division.id,
                    "division_nombre": division.nombre,
                    "componente_id": None,
                    "componente_nombre": "Valor",
                    "label": f"{division.nombre} - Valor" if division.nombre else resultado.nombre,
                    "unidad": resultado.unidad_medida or "",
                    "tipo_dato": "decimal",
                }, resultado=resultado, division=division))
                continue

            for componente in componentes:
                componente_label = componente.acronimo or componente.nombre
                if division.es_principal and (division.nombre or "").lower() in ["principal", "resultado simple"]:
                    label = componente_label
                else:
                    label = f"{division.nombre} - {componente_label}" if division.nombre else componente_label

                fields.append(enrich_with_limit({
                    "key": f"resultado:{resultado.id}:division:{division.id}:componente:{componente.id}",
                    "resultado_id": resultado.id,
                    "resultado_nombre": resultado.nombre,
                    "division_id": division.id,
                    "division_nombre": division.nombre,
                    "componente_id": componente.id,
                    "componente_nombre": componente_label,
                    "label": label,
                    "unidad": componente.division.resultado.unidad_medida or "",
                    "tipo_dato": componente.tipo_dato or "decimal",
                    "opciones": component_options(componente),
                    "escala_comparacion": componente.escala_comparacion_id,
                    "requiere_valor": componente.requiere_valor,
                    "permite_observacion": componente.permite_observacion,
                }, resultado=resultado, division=division, componente=componente))

    if not fields:
        fields.append(enrich_with_limit({
            "key": f"prueba:{prueba.id}:valor",
            "resultado_id": None,
            "resultado_nombre": prueba.nombre_variable,
            "division_id": None,
            "division_nombre": None,
            "componente_id": None,
            "componente_nombre": "Valor",
            "label": prueba.nombre_variable,
            "unidad": prueba.unidad_medida or prueba_muestra.unidad_configurada or "",
            "tipo_dato": "decimal",
        }))

    return fields


def _validate_result_value(field, value):
    if value in [None, ""]:
        return None

    tipo = field.get("tipo_dato") or "numerico"
    if tipo in ["numerico", "decimal"]:
        try:
            Decimal(str(value).replace(",", "."))
        except (InvalidOperation, TypeError, ValueError):
            return "Ingrese un valor numÃ©rico vÃ¡lido."

    if tipo == "booleano":
        normalized = str(value).strip().casefold()
        allowed = {"si", "sÃ­", "sí", "no", "true", "false", "1", "0", "pasa", "no pasa", "hay", "no hay"}
        for option in field.get("opciones") or []:
            allowed.add(str(option.get("value", "")).strip().casefold())
            allowed.add(str(option.get("label", "")).strip().casefold())
        if normalized not in allowed:
            return "Seleccione un valor booleano vÃ¡lido."

    if tipo in ["escala", "escala_ordinal"]:
        allowed = {str(option.get("value", option)).strip() for option in field.get("opciones") or []}
        if allowed and str(value).strip() not in allowed:
            return "Seleccione un valor permitido para este campo."

    if tipo in ["opcion", "texto"] and field.get("opciones"):
        allowed = {str(option.get("value", option)).strip() for option in field.get("opciones") or []}
        allowed.update(str(option.get("label", "")).strip() for option in field.get("opciones") or [])
        if str(value).strip() not in allowed:
            return "Seleccione un valor permitido para este campo."

    return None


def _normalize_imported_result_value(field, value):
    """Convierte valores de Excel al valor canonico usado por los controles web."""
    if value in [None, ""]:
        return value

    normalized = str(value).strip().casefold()
    for option in field.get("opciones") or []:
        option_value = str(option.get("value", "")).strip()
        option_label = str(option.get("label", "")).strip()
        if normalized in {option_value.casefold(), option_label.casefold()}:
            return option_value

    if (field.get("tipo_dato") or "").lower() == "booleano":
        if isinstance(value, bool):
            return "true" if value else "false"
        truthy = {"si", "sí", "true", "1", "pasa", "hay", "verdadero"}
        falsy = {"no", "false", "0", "no pasa", "no hay", "falso"}
        if normalized in truthy:
            return "true"
        if normalized in falsy:
            return "false"

    return value


def _serialize_entry_batch(lote):
    pruebas_qs = PruebaMuestra.objects.filter(muestra__lote=lote)
    total_pruebas = pruebas_qs.count()
    completadas = pruebas_qs.filter(estatus__in=["completado", "aprobado"]).count()
    borradores = pruebas_qs.filter(estatus="en_proceso").count()

    return {
        "id": lote.id,
        "cliente_nombre": getattr(lote.cliente_empresa, "name", None) or getattr(lote.cliente_empresa, "nombre", None) or lote.cliente_ocasional_nombre,
        "tipo_cliente": lote.tipo_cliente,
        "tipo_gestion": lote.tipo_gestion_id,
        "tipo_gestion_nombre": getattr(lote.tipo_gestion, "nombre", None),
        "fecha_recepcion": lote.fecha_recepcion,
        "estado": lote.estado,
        "total_muestras": lote.total_muestras,
        "total_pruebas": total_pruebas,
        "pruebas_completadas": completadas,
        "pruebas_borrador": borradores,
        "pruebas_pendientes": max(total_pruebas - completadas, 0),
    }


def _result_payload(result):
    if not result:
        return None
    payload = _safe_json_loads(result.resultado)
    return {
        "id": result.id,
        "estatus": result.estatus,
        "resultado": result.resultado,
        "payload": payload,
        "observaciones": result.observaciones,
        "fecha_medicion": result.fecha_medicion,
        "fecha_actualizacion": result.fecha_actualizacion,
        "usuario_medicion": _user_payload(result.usuario_medicion),
    }


def _flatten_payload_values(payload):
    if not isinstance(payload, dict):
        return []
    fields = payload.get("fields") or []
    values = payload.get("values") or []
    field_map = {field.get("key"): field for field in fields if field.get("key")}
    flattened = []
    for item in values:
        key = item.get("key")
        field = field_map.get(key, {})
        value = item.get("value")
        value_label = item.get("value_label") or _display_value_for_field(field, value)
        flattened.append({
            "key": key,
            "label": field.get("label") or field.get("componente_nombre") or key,
            "unidad": field.get("unidad") or "",
            "value": value,
            "value_label": value_label,
            "observacion": item.get("observacion") or "",
            "resultado_nombre": field.get("resultado_nombre"),
            "division_nombre": field.get("division_nombre"),
            "componente_nombre": field.get("componente_nombre"),
        })
    return flattened


def _result_summary(result, prueba_muestra=None):
    payload = _safe_json_loads(result.resultado) if result else None
    flat = _flatten_payload_values(payload)
    if flat:
        clean = [item for item in flat if item.get("value") not in [None, ""]]
        if len(clean) == 1:
            item = clean[0]
            unit = item.get("unidad") or ""
            return f"{item.get('value_label') or item.get('value')} {unit}".strip()

        # Para pruebas compuestas por divisiones, por ejemplo Espuma:
        # Secuencia I: EI/FI, Secuencia II: EII/FII, Secuencia III: EIII/FIII
        groups = []
        current_key = object()
        current_values = []
        for item in clean:
            key = item.get("division_nombre") or item.get("resultado_nombre") or "__default__"
            if current_values and key != current_key:
                groups.append(" / ".join(str(value) for value in current_values))
                current_values = []
            current_key = key
            current_values.append(item.get("value_label") or item.get("value"))
        if current_values:
            groups.append(" / ".join(str(value) for value in current_values))

        if groups:
            return " · ".join(groups)

    if prueba_muestra and prueba_muestra.valor:
        return prueba_muestra.valor
    return "-"


def _sample_equipment_label(muestra):
    if getattr(muestra, "referencia_equipo_id", None) and getattr(muestra, "referencia_equipo", None):
        return getattr(muestra.referencia_equipo, "nombre", None) or getattr(muestra.referencia_equipo, "name", None) or str(muestra.referencia_equipo)
    return muestra.equipo_placa or "-"


def _sample_review_status(muestra, tests):
    if not tests:
        return "Sin pruebas"
    statuses = {
        _canonical_status(test.get("estado_tecnico") or test.get("estado_limite"))
        for test in tests
    }
    if "CRITICO" in statuses:
        return "Crítica"
    if "FUERA_DE_LIMITE" in statuses:
        return "Con alerta"
    if all(t.get("is_revisada") for t in tests):
        return "Revisada"
    if any(t.get("resultado") for t in tests):
        return "Pendiente"
    return "Sin resultado"



def _operator_symbol(operator):
    return {
        "max": "<=",
        "warn_max": "<=",
        "min": ">=",
        "warn_min": ">=",
        "eq": "=",
    }.get(operator or "", operator or "")


def _clean_limit_field_label(value):
    label = str(value or "").strip()
    if not label:
        return ""
    replacements = {
        "secuencia_1_": "Sec. I ",
        "secuencia_2_": "Sec. II ",
        "secuencia_3_": "Sec. III ",
        "_valor_minimo": "",
        "_valor_maximo": "",
        "_min": "",
        "_max": "",
    }
    for source, target in replacements.items():
        label = label.replace(source, target)
    label = label.replace("_", " ").strip()
    compact = {
        "ei": "EI", "fi": "FI", "eii": "EII", "fii": "FII", "eiii": "EIII", "fiii": "FIII",
        "4um": "4um", "6um": "6um", "14um": "14um",
    }
    return compact.get(label.lower(), label)


def _limit_detail_label(detail):
    field = _clean_limit_field_label(
        detail.get("field_label") or detail.get("field") or detail.get("componente") or detail.get("resultado") or ""
    )
    raw_limit_value = detail.get("limit_value")
    # Preserve semaphore rules as structured data. Their display label is a
    # Python-style dictionary string and must never leak into API summaries.
    limit_value = raw_limit_value if isinstance(raw_limit_value, dict) else (
        detail.get("limit_value_label") or raw_limit_value
    )
    if isinstance(limit_value, str) and limit_value.strip().startswith("{"):
        try:
            limit_value = json.loads(limit_value)
        except (TypeError, ValueError):
            pass
    if isinstance(limit_value, dict):
        parts = []
        if limit_value.get("min_aceptable") not in [None, ""]:
            parts.append(f"aceptable desde {limit_value['min_aceptable']}")
        if limit_value.get("max_aceptable") not in [None, ""]:
            parts.append(f"aceptable hasta {limit_value['max_aceptable']}")
        if limit_value.get("min_critico") not in [None, ""]:
            parts.append(f"critico bajo {limit_value['min_critico']}")
        if limit_value.get("max_critico") not in [None, ""]:
            parts.append(f"critico sobre {limit_value['max_critico']}")
        if limit_value.get("valor_esperado") not in [None, ""]:
            parts.append(f"esperado {limit_value['valor_esperado']}")
        label = " · ".join(parts) or "Regla configurada"
        return f"{field}: {label}" if field else label
    unit = detail.get("unit") or ""
    operator = _operator_symbol(detail.get("operator"))
    if limit_value in [None, ""]:
        return field or detail.get("estado") or "Sin límite"
    label = f"{operator} {limit_value} {unit}".strip()
    return f"{field}: {label}" if field else label


def _limit_summary_from_evaluation(evaluation):
    if not evaluation:
        return "No evaluable"
    if hasattr(evaluation, "to_dict"):
        data = evaluation.to_dict()
    else:
        data = evaluation

    if "results" in data:
        results = data.get("results") or []
        comparable = [item for item in results if item.get("limit_value") not in [None, ""]]
        if comparable:
            return ", ".join(_limit_detail_label(item) for item in comparable[:6])
        return data.get("reason") or data.get("estado") or "Sin límite configurado"

    if data.get("limit_value") not in [None, ""]:
        return _limit_detail_label(data)
    return data.get("reason") or data.get("estado") or "Sin límite configurado"


def _evaluation_status_label(estado):
    mapping = {
        "NORMAL": "NORMAL",
        "FUERA_DE_LIMITE": "NO DESEADO",
        "NO_DESEADO": "NO DESEADO",
        "CRITICO": "CRÍTICO",
        "SIN_INFORMACION_TECNICA": "SIN INFO TÉCNICA",
        "SIN_LIMITE_CONFIGURADO": "SIN LÍMITE CONFIGURADO",
        "SIN_LIMITE": "SIN LÍMITE",
        "SIN_RESULTADO": "SIN RESULTADO",
        "SIN_CRITERIO": "SIN CRITERIO",
        "NO_EVALUABLE": "NO EVALUABLE",
        "INFORMATIVO": "INFORMATIVO",
    }
    return mapping.get(estado or "", estado or "NO EVALUABLE")


def _canonical_status(value):
    label = str(value or "").strip().upper().replace(" ", "_").replace("-", "_")
    aliases = {
        "NO_DESEADO": "FUERA_DE_LIMITE",
        "CRITICO": "CRITICO",
        "FUERA_DE_LIMITE": "FUERA_DE_LIMITE",
        "FUERA_LIMITE": "FUERA_DE_LIMITE",
        "ALERTA": "FUERA_DE_LIMITE",
        "NORMAL": "NORMAL",
        "OK": "NORMAL",
        "CUMPLE": "NORMAL",
        "INFORMATIVO": "INFORMATIVO",
        "SIN_LIMITE": "SIN_LIMITE",
        "SIN_LÍMITE": "SIN_LIMITE",
        "SIN_LIMITE_CONFIGURADO": "SIN_LIMITE_CONFIGURADO",
        "SIN_INFORMACION_TECNICA": "SIN_INFORMACION_TECNICA",
        "SIN_INFO_TECNICA": "SIN_INFORMACION_TECNICA",
        "SIN_RESULTADO": "SIN_RESULTADO",
        "NO_EVALUABLE": "NO_EVALUABLE",
    }
    return aliases.get(label, label or "NO_EVALUABLE")


def _is_bad_or_incomplete_evaluation(evaluation):
    if not evaluation:
        return True
    estado = _canonical_status(evaluation.get("estado") or evaluation.get("estado_label"))
    if estado in ["NORMAL", "FUERA_DE_LIMITE", "INFORMATIVO"]:
        return False
    # If there are useful comparable details, keep it even if aggregate label is imperfect.
    for detail in evaluation.get("details") or []:
        if detail.get("limit_value") not in [None, ""] and detail.get("passed") is not None:
            return False
    return True


def _field_is_informative(field):
    tipo = str(field.get("tipo_dato") or field.get("tipo_comparacion") or "").lower()
    if tipo in ["comentario", "informativo"]:
        return True
    if field.get("requiere_valor") is False and tipo not in ["numerico", "decimal", "booleano", "escala", "escala_ordinal", "opcion"]:
        return True
    return False


def _field_value_key(field):
    return str(field.get("key") or "")


def _norm_lookup_text(value):
    return str(value or "").strip().casefold().replace("_", " ")


def _entry_label(entry):
    field = entry.get("field") or {}
    return field.get("label") or field.get("componente_nombre") or field.get("division_nombre") or field.get("resultado_nombre") or entry.get("key") or ""


def _detail_lookup_keys(detail):
    values = [
        detail.get("field_id"),
        detail.get("field"),
        detail.get("field_label"),
        detail.get("componente"),
        detail.get("resultado"),
    ]
    return {_norm_lookup_text(v) for v in values if v not in [None, ""]}


def _entry_lookup_keys(entry):
    field = entry.get("field") or {}
    values = [
        field.get("limite_campo"),
        field.get("codigo"),
        field.get("key"),
        field.get("label"),
        field.get("componente_nombre"),
        field.get("resultado_nombre"),
        entry.get("key"),
    ]
    return {_norm_lookup_text(v) for v in values if v not in [None, ""]}


def _enrich_evaluation_details(evaluation, entries):
    if not evaluation:
        return evaluation
    entry_pool = [entry for entry in entries if entry.get("valor") not in [None, ""]]
    enriched = []
    used = set()

    for idx, detail in enumerate(evaluation.get("details") or []):
        detail = dict(detail)
        matched = None
        dkeys = _detail_lookup_keys(detail)
        for j, entry in enumerate(entry_pool):
            if j in used:
                continue
            if dkeys & _entry_lookup_keys(entry):
                matched = entry
                used.add(j)
                break
        if matched is None and idx < len(entry_pool):
            matched = entry_pool[idx]
            used.add(idx)
        if matched:
            field = matched.get("field") or {}
            detail.setdefault("value", matched.get("valor"))
            detail.setdefault("resultado_valor", matched.get("valor"))
            detail.setdefault("resultado_valor_label", _display_value_for_field(field, matched.get("valor")))
            detail.setdefault("limit_value_label", _display_value_for_field(field, detail.get("limit_value")))
            detail.setdefault("result_label", field.get("resultado_nombre"))
            detail.setdefault("division_id", field.get("division_id"))
            detail.setdefault("division_label", field.get("division_nombre"))
            detail.setdefault("component_id", field.get("componente_id"))
            detail.setdefault("field_label", field.get("label") or field.get("componente_nombre") or detail.get("field_label"))
            detail.setdefault("unit", detail.get("unit") or field.get("unidad"))
            detail.setdefault("tipo_dato", field.get("tipo_dato") or field.get("tipo_comparacion"))
        enriched.append(detail)

    evaluation["details"] = enriched
    evaluation["estado"] = _canonical_status(evaluation.get("estado"))
    evaluation["estado_label"] = _evaluation_status_label(evaluation.get("estado"))
    if not evaluation.get("summary"):
        evaluation["summary"] = _limit_summary_from_evaluation({"results": enriched, "estado": evaluation.get("estado")})
    return evaluation


def _criteria_summary_from_test(prueba_muestra, evaluation=None):
    parts = []
    if getattr(prueba_muestra, "criterio_limite_catalogo_id", None) and getattr(prueba_muestra, "criterio_limite_item_id", None):
        catalog_name = getattr(prueba_muestra.criterio_limite_catalogo, "nombre", None) or getattr(prueba_muestra.criterio_limite_catalogo, "codigo", None)
        item_name = getattr(prueba_muestra.criterio_limite_item, "nombre", None) or getattr(prueba_muestra.criterio_limite_item, "codigo", None)
        if catalog_name and item_name:
            parts.append(f"{catalog_name}: {item_name}")
    if getattr(prueba_muestra, "criterio_limite_escala_id", None) and getattr(prueba_muestra, "criterio_limite_escala_item_id", None):
        escala = getattr(prueba_muestra.criterio_limite_escala, "nombre", None) or getattr(prueba_muestra.criterio_limite_escala, "codigo", None)
        item = getattr(prueba_muestra.criterio_limite_escala_item, "etiqueta", None)
        if escala and item:
            parts.append(f"{escala}: {item}")
    if getattr(prueba_muestra, "criterio_limite_valor", None) not in [None, ""]:
        raw_value = _safe_json_loads(prueba_muestra.criterio_limite_valor)
        if isinstance(raw_value, dict):
            fields = {
                str(field.id): field
                for field in PruebaLimiteCampo.objects.filter(
                    fuente_limite__prueba=prueba_muestra.prueba,
                    activo=True,
                    deleted_at__isnull=True,
                ).select_related("componente")
            }
            field_parts = []
            for key, value in raw_value.items():
                field = fields.get(str(key))
                if not field:
                    continue
                display = _display_value_for_field({
                    "opciones": [
                        {"value": "true", "label": getattr(field.componente, "etiqueta_verdadero", None) or "Sí"},
                        {"value": "false", "label": getattr(field.componente, "etiqueta_falso", None) or "No"},
                    ] if field.tipo_comparacion == "booleano" else []
                }, value)
                field_parts.append(f"{field.nombre}: {display}")
            parts.append(" / ".join(field_parts) if field_parts else f"Valor asignado: {prueba_muestra.criterio_limite_valor}")
        else:
            parts.append(f"Valor asignado: {prueba_muestra.criterio_limite_valor}")

    for detail in (evaluation or {}).get("details") or []:
        catalog_label = detail.get("catalog_label") or detail.get("catalog")
        item_label = detail.get("item_label")
        if catalog_label and item_label:
            text = f"{catalog_label}: {item_label}"
            if text not in parts:
                parts.append(text)
        resolution = detail.get("resolution_reason")
        if resolution and not parts:
            parts.append(resolution)
    if parts:
        return " · ".join(parts[:3])
    estado = _canonical_status((evaluation or {}).get("estado"))
    if estado == "SIN_LIMITE":
        return "No aplica"
    if estado == "INFORMATIVO":
        return "Informativo"
    if estado == "SIN_INFORMACION_TECNICA":
        return "Sin información técnica"
    if estado == "SIN_LIMITE_CONFIGURADO":
        return "Sin límite configurado"
    return "Definido en límites"



def _compare_limit_value(value, operator, limit_value):
    try:
        value_dec = Decimal(str(value).replace(",", "."))
        limit_dec = Decimal(str(limit_value).replace(",", "."))
    except (InvalidOperation, TypeError, ValueError):
        return str(value).strip().casefold() == str(limit_value).strip().casefold() if operator == "eq" else None

    if operator in ["max", "warn_max"]:
        return value_dec <= limit_dec
    if operator in ["min", "warn_min"]:
        return value_dec >= limit_dec
    if operator == "eq":
        return value_dec == limit_dec
    return None


def _field_matches_component(field, componente):
    if not componente:
        return False
    token = (getattr(componente, "acronimo", None) or getattr(componente, "nombre", None) or "").strip().lower()
    if not token:
        return False
    candidates = [
        str(field.codigo or "").lower(),
        str(field.nombre or "").lower(),
    ]
    for candidate in candidates:
        normalized = candidate.replace("-", "_").replace(" ", "_")
        parts = [part for part in normalized.split("_") if part]
        if token in parts or normalized.endswith(f"_{token}") or normalized == token:
            return True
    return False


def _matching_limit_fields(source, resultado=None, division=None, componente=None):
    queryset = PruebaLimiteCampo.objects.filter(
        fuente_limite=source,
        activo=True,
        deleted_at__isnull=True,
    ).select_related("resultado", "division", "componente").order_by("orden", "id")

    if componente:
        exact = list(queryset.filter(componente=componente))
        if exact:
            return exact

    if division:
        division_fields = list(queryset.filter(componente__isnull=True, division=division))
        if division_fields:
            matched = [field for field in division_fields if _field_matches_component(field, componente)]
            return matched or division_fields

    if resultado:
        result_fields = list(queryset.filter(
            componente__isnull=True,
            division__isnull=True,
            resultado=resultado,
        ))
        if result_fields:
            matched = [field for field in result_fields if _field_matches_component(field, componente)]
            return matched or result_fields

    global_fields = list(queryset.filter(
        componente__isnull=True,
        division__isnull=True,
        resultado__isnull=True,
    ))
    if componente:
        matched = [field for field in global_fields if _field_matches_component(field, componente)]
        if matched:
            return matched
    return global_fields


def _evaluate_sample_test_limits(prueba_muestra, result=None):
    """
    Evalúa límites usando un contrato único para revisión e interpretación.

    Salida estable:
    - estado / estado_label
    - summary
    - details[] con resultado_valor, field_label, operator, limit_value, estado, passed
    - raw
    """
    result = result or _active_result_for(prueba_muestra)
    if not result:
        return {
            "estado": "SIN_RESULTADO",
            "estado_label": "SIN RESULTADO",
            "summary": "Sin resultado",
            "details": [],
            "raw": None,
        }

    payload = _safe_json_loads(result.resultado)
    if not isinstance(payload, dict):
        value = prueba_muestra.valor
        evaluation = resolve_and_evaluate_result(
            value,
            prueba_muestra.muestra,
            prueba_muestra.prueba,
        )
        data = evaluation.to_dict()
        data = {
            "estado": _canonical_status(data.get("estado")),
            "estado_label": _evaluation_status_label(data.get("estado")),
            "summary": _limit_summary_from_evaluation(data),
            "details": [data],
            "raw": data,
        }
        return _enrich_evaluation_details(data, [{"valor": value, "field": {"label": prueba_muestra.prueba.nombre_variable}}])

    fields = payload.get("fields") or []
    values = payload.get("values") or []
    current_fields = _build_result_fields(prueba_muestra)
    field_map = {field.get("key"): field for field in fields if field.get("key")}
    current_field_map = {field.get("key"): field for field in current_fields if field.get("key")}

    entries = []
    evaluable = []
    informative_entries = []
    for index, item in enumerate(values):
        key = item.get("key")
        value = item.get("value")
        if value in [None, ""]:
            continue
        # Los resultados guardados antes no traían payload.fields. Rehidratar
        # desde la estructura actual de la prueba permite evaluar compuestas
        # como espuma/nivel de contaminación sin reingresar resultados.
        field = field_map.get(key) or current_field_map.get(key) or (
            current_fields[index] if index < len(current_fields) else {}
        )

        resultado = PruebaResultado.objects.filter(pk=field.get("resultado_id")).first() if field.get("resultado_id") else None
        division = PruebaResultadoDivision.objects.filter(pk=field.get("division_id")).first() if field.get("division_id") else None
        componente = PruebaResultadoComponente.objects.filter(pk=field.get("componente_id")).first() if field.get("componente_id") else None

        entry = {
            "valor": value,
            "resultado": resultado,
            "division": division,
            "componente": componente,
            "field": field,
            "key": key,
        }
        entries.append(entry)
        if _field_is_informative(field) or getattr(componente, "tipo_dato", None) == "comentario":
            informative_entries.append(entry)
        else:
            evaluable.append(entry)

    if not entries:
        return {
            "estado": "SIN_RESULTADO",
            "estado_label": "SIN RESULTADO",
            "summary": "Sin resultado capturado",
            "details": [],
            "raw": None,
        }

    if not evaluable:
        details = [
            {
                "applies": False,
                "estado": "INFORMATIVO",
                "field_label": _entry_label(entry),
                "resultado_valor": entry.get("valor"),
                "value": entry.get("valor"),
                "limit_value": None,
                "operator": None,
                "unit": (entry.get("field") or {}).get("unidad") or "",
                "passed": None,
                "reason": "Campo informativo sin evaluación automática.",
            }
            for entry in informative_entries
        ]
        return {
            "estado": "INFORMATIVO",
            "estado_label": "INFORMATIVO",
            "summary": "No aplica",
            "details": details,
            "raw": {"estado": "INFORMATIVO", "results": details},
        }

    # 1) Motor principal de límites.
    evaluation = resolve_and_evaluate_values_for_sample_test(prueba_muestra, evaluable)
    data = evaluation.to_dict()
    primary = {
        "estado": _canonical_status(data.get("estado")),
        "estado_label": _evaluation_status_label(data.get("estado")),
        "summary": _limit_summary_from_evaluation(data),
        "details": data.get("results") or [],
        "reason": data.get("reason") or "",
        "result_outcomes": data.get("result_outcomes") or [],
        "raw": data,
    }
    primary = _enrich_evaluation_details(primary, evaluable)

    primary["estado"] = _canonical_status(primary.get("estado"))
    primary["estado_label"] = _evaluation_status_label(primary.get("estado"))
    primary["summary"] = _limit_summary_from_evaluation({"results": primary.get("details") or [], "estado": primary.get("estado")})
    return primary


def _sync_sample_test_limit_evaluation(prueba_muestra, result=None):
    evaluation = _evaluate_sample_test_limits(prueba_muestra, result)
    prueba_muestra.evaluacion_limite = evaluation
    prueba_muestra.estado_limite = evaluation.get("estado_label") or evaluation.get("estado")
    prueba_muestra.save(update_fields=["evaluacion_limite", "estado_limite"])
    return evaluation


def _test_structure_payload(prueba):
    """Serialize the authored result layout used by every operational flow."""
    results = []
    active_results = sorted(
        (item for item in prueba.resultados.all() if item.activo),
        key=lambda item: (item.orden, item.id),
    )
    for result in active_results:
        divisions = []
        active_divisions = sorted(
            (item for item in result.divisiones.all() if item.activo),
            key=lambda item: (item.orden, item.id),
        )
        for division in active_divisions:
            components = []
            active_components = sorted(
                (item for item in division.componentes.all() if item.activo),
                key=lambda item: (item.orden, item.id),
            )
            for component in active_components:
                scale = component.escala_comparacion
                scale_items = []
                if scale:
                    scale_items = sorted(
                        (
                            item for item in scale.items.all()
                            if item.activo and item.deleted_at is None
                        ),
                        key=lambda item: (item.orden, item.id),
                    )
                components.append({
                    "id": component.id,
                    "nombre": component.nombre,
                    "acronimo": component.acronimo,
                    "tipo_dato": component.tipo_dato,
                    "orden": component.orden,
                    "requiere_valor": component.requiere_valor,
                    "permite_observacion": component.permite_observacion,
                    "etiqueta_verdadero": component.etiqueta_verdadero,
                    "etiqueta_falso": component.etiqueta_falso,
                    "opciones_resultado": component.opciones_resultado or [],
                    "escala_comparacion": ({
                        "id": scale.id,
                        "nombre": scale.nombre,
                        "codigo": scale.codigo,
                        "items": [
                            {
                                "id": scale_item.id,
                                "etiqueta": scale_item.etiqueta,
                                "valor_normalizado": scale_item.valor_normalizado,
                                "orden": scale_item.orden,
                            }
                            for scale_item in scale_items
                        ],
                    } if scale else None),
                })

            separators = [
                {"id": separator.id, "simbolo": separator.simbolo, "orden": separator.orden}
                for separator in sorted(
                    (item for item in division.separadores.all() if item.activo),
                    key=lambda item: (item.orden, item.id),
                )
            ]
            dispositions = sorted(
                (item for item in division.disposiciones.all() if item.activo),
                key=lambda item: (item.orden, item.id),
            )
            disposition = dispositions[0] if dispositions else None
            disposition_items = []
            if disposition:
                for disposition_item in sorted(
                    disposition.items.all(),
                    key=lambda item: (item.orden, item.id),
                ):
                    if disposition_item.tipo == "componente" and disposition_item.componente:
                        value = disposition_item.componente.acronimo or disposition_item.componente.nombre
                    elif disposition_item.separador:
                        value = disposition_item.separador.simbolo
                    else:
                        value = ""
                    disposition_items.append({
                        "id": disposition_item.id,
                        "tipo": disposition_item.tipo,
                        "componente_id": disposition_item.componente_id,
                        "separador_id": disposition_item.separador_id,
                        "valor": value,
                        "orden": disposition_item.orden,
                    })

            divisions.append({
                "id": division.id,
                "nombre": division.nombre,
                "es_principal": division.es_principal,
                "orden": division.orden,
                "componentes": components,
                "separadores": separators,
                "disposicion": disposition_items,
            })
        results.append({
            "id": result.id,
            "nombre": result.nombre,
            "acronimo": result.acronimo,
            "descripcion": result.descripcion,
            "unidad_medida": result.unidad_medida,
            "unidad_catalogo": result.unidad_catalogo_id,
            "orden": result.orden,
            "divisiones": divisions,
        })
    return results


def _sample_tests_for_lote(lote, muestra_id=None, include_test_structure=True):
    queryset = (
        PruebaMuestra.objects
        .filter(muestra__lote=lote, estado_asignacion="confirmada")
        .select_related(
            "muestra",
            "muestra__referencia_equipo",
            "prueba",
            "prueba__metodo",
            "prueba__metodo__equipo_prueba",
            "equipo_configurado",
            "metodo_configurado",
            "criterio_limite_catalogo",
            "criterio_limite_item",
            "criterio_limite_item__catalogo",
            "criterio_limite_escala",
            "criterio_limite_escala_item",
            "criterio_limite_escala_item__escala",
            "usuario_solicitud",
            "usuario_medicion",
            "usuario_revision",
        )
        .prefetch_related(
            "resultados",
            "prueba__resultados__divisiones__componentes__escala_comparacion__items",
            "prueba__resultados__divisiones__separadores",
            "prueba__resultados__divisiones__disposiciones__items__componente",
            "prueba__resultados__divisiones__disposiciones__items__separador",
        )
        .order_by("muestra_id", "prueba__nombre_variable")
    )
    if muestra_id is not None:
        queryset = queryset.filter(muestra_id=muestra_id)
    rows = []
    for st in queryset:
        result = _active_result_for(st)
        limit_evaluation = _evaluate_sample_test_limits(st, result)
        technical_status = _canonical_status(limit_evaluation.get("estado"))
        technical_label = _evaluation_status_label(technical_status)
        criterion_summary = _criteria_summary_from_test(st, limit_evaluation)
        rows.append({
            "id": st.id,
            "muestra": st.muestra_id,
            "prueba": {
                "id": st.prueba_id,
                "nombre_variable": st.prueba.nombre_variable,
                "acronimo": st.prueba.acronimo,
                "unidad_medida": st.prueba.unidad_medida,
                "condicion": st.condicion_configurada or st.prueba.condicion,
                # La estructura completa se usa en revisión, pero en interpretación
                # ya vienen el resultado y la evaluación resueltos. Omitirla allí
                # evita repetir decenas de KB por prueba.
                "resultados": _test_structure_payload(st.prueba) if include_test_structure else [],
            },
            "equipo": (
                getattr(st.equipo_configurado, "nombre", None)
                or getattr(getattr(st.metodo_configurado, "equipo_prueba", None), "nombre", None)
                or getattr(getattr(getattr(st.prueba, "metodo", None), "equipo_prueba", None), "nombre", None)
                or "-"
            ),
            "equipo_codigo": (
                getattr(st.equipo_configurado, "codigo", None)
                or getattr(getattr(st.metodo_configurado, "equipo_prueba", None), "codigo", None)
                or getattr(getattr(getattr(st.prueba, "metodo", None), "equipo_prueba", None), "codigo", None)
                or "-"
            ),
            "metodo": getattr(st.metodo_configurado, "nombre", None) or getattr(getattr(st.prueba, "metodo", None), "nombre", None) or "-",
            "metodo_codigo": getattr(st.metodo_configurado, "codigo", None) or getattr(getattr(st.prueba, "metodo", None), "codigo", None) or "-",
            "resultado": _result_payload(result),
            "resultado_resumen": _result_summary(result, st),
            "fecha_medicion": getattr(result, "fecha_medicion", None) or st.fecha_medicion,
            "usuario_medicion": _user_payload(st.usuario_medicion) or (_result_payload(result).get("usuario_medicion") if result and _result_payload(result) else None),
            "estatus": st.estatus,
            "is_revisada": bool(st.is_revisada),
            "estado_tecnico": technical_status,
            "estado_tecnico_label": technical_label,
            "estado_limite": technical_label,
            "limite_resumen": limit_evaluation.get("summary"),
            "limite_detalles": limit_evaluation.get("details") or [],
            "criterio_aplicado": criterion_summary,
            "evaluacion_limite": limit_evaluation,
            "criterio_limite_catalogo": st.criterio_limite_catalogo_id,
            "criterio_limite_item": st.criterio_limite_item_id,
            "criterio_limite_catalogo_info": {
                "id": st.criterio_limite_catalogo_id,
                "nombre": st.criterio_limite_catalogo.nombre,
                "codigo": st.criterio_limite_catalogo.codigo,
            } if st.criterio_limite_catalogo_id else None,
            "criterio_limite_item_info": {
                "id": st.criterio_limite_item_id,
                "nombre": st.criterio_limite_item.nombre,
                "codigo": st.criterio_limite_item.codigo,
                "catalogo": st.criterio_limite_item.catalogo_id,
            } if st.criterio_limite_item_id else None,
            "criterio_limite_escala": st.criterio_limite_escala_id,
            "criterio_limite_escala_item": st.criterio_limite_escala_item_id,
            "criterio_limite_valor": st.criterio_limite_valor,
            "criterio_limite_escala_info": {
                "id": st.criterio_limite_escala_id,
                "nombre": st.criterio_limite_escala.nombre,
                "codigo": st.criterio_limite_escala.codigo,
            } if st.criterio_limite_escala_id else None,
            "criterio_limite_escala_item_info": {
                "id": st.criterio_limite_escala_item_id,
                "etiqueta": st.criterio_limite_escala_item.etiqueta,
                "orden": st.criterio_limite_escala_item.orden,
                "escala": st.criterio_limite_escala_item.escala_id,
            } if st.criterio_limite_escala_item_id else None,
            "unidad": st.unidad_configurada or st.unidad or st.prueba.unidad_medida,
            "configuracion_resultados": st.configuracion_resultados or {},
            "observaciones": st.observaciones,
        })
    return rows


def _batch_review_payload(lote):
    tests = _sample_tests_for_lote(lote)
    samples = []
    for muestra in lote.muestras.select_related("referencia_equipo").all().order_by("id"):
        sample_tests = [item for item in tests if item["muestra"] == muestra.id]
        interpretation = _get_interpretation_data(muestra)
        has_interpretation = bool(
            interpretation.get("conclusion")
            or interpretation.get("comentarios_predefinidos")
            or interpretation.get("comentarios_manuales")
        )
        samples.append({
            "id": muestra.id,
            "fecha_toma": muestra.fecha_toma,
            "tipo_muestra": muestra.tipo_muestra,
            "condicion": muestra.condicion,
            "referencia_marca": muestra.referencia_marca,
            "equipo": _sample_equipment_label(muestra),
            "ubicacion": getattr(getattr(muestra, "referencia_equipo", None), "location", None) or getattr(getattr(muestra, "referencia_equipo", None), "ubicacion", None) or muestra.equipo_placa or "-",
            "is_revisado": muestra.is_revisado,
            "is_resultado_ingresado": muestra.is_resultado_ingresado,
            "estado_revision": _sample_review_status(muestra, sample_tests),
            "tests_count": len(sample_tests),
            "reviewed_count": len([item for item in sample_tests if item.get("is_revisada")]),
            "has_interpretation": has_interpretation,
        })
    return {
        "id": lote.id,
        "cliente_nombre": getattr(lote.cliente_empresa, "nombre", None) or getattr(lote.cliente_empresa, "name", None) or lote.cliente_ocasional_nombre,
        "tipo_gestion": lote.tipo_gestion_id,
        "tipo_gestion_nombre": getattr(lote.tipo_gestion, "nombre", None),
        "estado": lote.estado,
        "fecha_recepcion": lote.fecha_recepcion,
        "total_muestras": lote.total_muestras,
        "muestras_revisadas": lote.muestras_revisadas,
        "muestras": samples,
        "sample_tests": tests,
    }


def _build_predefined_comments(sample_tests):
    comments = []
    for item in sample_tests:
        prueba = item.get("prueba", {})
        nombre = (prueba.get("nombre_variable") or prueba.get("acronimo") or "").lower()
        result = item.get("resultado_resumen") or "-"
        estado = (item.get("estado_limite") or "").lower()
        if "viscos" in nombre:
            comments.append("La viscosidad del aceite se encuentra dentro del rango normal para la condición evaluada.")
        elif "agua" in nombre:
            comments.append("El contenido de agua se encuentra dentro del rango esperado para la muestra analizada.")
        elif "contamin" in nombre or "part" in nombre or "iso" in nombre:
            comments.append("El nivel de contaminación debe revisarse frente a la tabla de limpieza aplicable.")
        elif estado in ["fuera", "alerta", "critico", "crítico"]:
            comments.append(f"El ensayo {prueba.get('acronimo') or prueba.get('nombre_variable')} requiere seguimiento por valor {result}.")
    # Deduplicate preserving order.
    unique = []
    for comment in comments:
        if comment not in unique:
            unique.append(comment)
    return unique[:8]


def _get_interpretation_data(muestra):
    extra = muestra.campos_adicionales or {}
    return extra.get("interpretacion", {}) if isinstance(extra, dict) else {}


def _set_interpretation_data(muestra, data):
    extra = muestra.campos_adicionales or {}
    if not isinstance(extra, dict):
        extra = {}
    extra["interpretacion"] = data
    muestra.campos_adicionales = extra
    muestra.save(update_fields=["campos_adicionales"])


class ResultadoViewSet(ActionPermissionMixin, viewsets.ModelViewSet):
    pagination_class = StandardResultsSetPagination
    permission_classes = [IsAuthenticated, DenyReadOnlyWrite, HasSecurityPermission]
    permission_action_map = {
        "list": "resultados.ver",
        "retrieve": "resultados.ver",
        "create": "resultados.cargar",
        "update": "resultados.editar",
        "partial_update": "resultados.editar",
        "destroy": "resultados.eliminar",
        "revisar": "revision.editar_estado",
        "entry_batches": "resultados.cargar",
        "entry_batch_detail": "resultados.cargar",
        "entry_form": "resultados.cargar",
        "save_draft": "resultados.cargar",
        "confirm_result": "resultados.finalizar",
        "confirm_batch_results": "resultados.finalizar",
        "excel_template": "resultados.importar",
        "excel_export": "resultados.exportar",
        "excel_preview": "resultados.importar",
        "excel_import": "resultados.importar",
        "review_batches": "revision.ver",
        "review_batch_detail": "revision.ver",
        "result_history": "resultados.ver_historico",
        "correct_result": "revision.editar_estado",
        "mark_sample_test_reviewed": "revision.aprobar_resultados",
        "complete_batch_review": "revision.aprobar_resultados",
        "interpretation_batch_detail": "interpretacion.ver",
        "interpretation_sample": "interpretacion.ver",
        "save_interpretation_sample": "interpretacion.guardar",
        "interpretation_trends": "interpretacion.ver_tendencias",
    }
    queryset = Resultado.objects.select_related(
        'prueba_muestra',
        'prueba_muestra__muestra',
        'prueba_muestra__muestra__lote',
        'prueba_muestra__prueba',
        'usuario_medicion',
    )

    def get_queryset(self):
        queryset = super().get_queryset()
        user = self.request.user
        if user.is_superuser or user.role == User.Role.GLOBAL:
            return queryset
        return queryset.filter(prueba_muestra__muestra__lote__cliente_empresa=user.empresa)

    def _scope_lotes(self, queryset=None):
        queryset = queryset or LoteMuestras.objects.all()
        user = self.request.user
        if user.is_superuser or user.role == User.Role.GLOBAL:
            return queryset
        return queryset.filter(cliente_empresa=user.empresa)

    def _scope_sample_tests(self, queryset=None):
        queryset = queryset or PruebaMuestra.objects.all()
        user = self.request.user
        if user.is_superuser or user.role == User.Role.GLOBAL:
            return queryset
        return queryset.filter(muestra__lote__cliente_empresa=user.empresa)

    def _scope_samples(self, queryset=None):
        queryset = queryset or Muestra.objects.all()
        user = self.request.user
        if user.is_superuser or user.role == User.Role.GLOBAL:
            return queryset
        return queryset.filter(lote__cliente_empresa=user.empresa)

    def get_serializer_class(self):
        if self.action in ['create', 'update', 'partial_update']:
            return CreateResultadoSerializer
        return ResultadoSerializer

    def perform_create(self, serializer):
        serializer.save(usuario_medicion=self.request.user)

    def _save_result_payload(
        self,
        request,
        prueba_muestra_id,
        estatus,
        *,
        defer_rollups=False,
        compact_response=False,
    ):
        prueba_muestra = get_object_or_404(
            PruebaMuestra.objects.select_related("muestra", "muestra__lote", "prueba"),
            pk=prueba_muestra_id,
        )
        require_batch_operation(prueba_muestra.muestra.lote, "ingresar_resultados")
        values = request.data.get("values", [])
        observaciones = request.data.get("observaciones", "")
        fecha_medicion = request.data.get("fecha_medicion")
        equipo_prueba = request.data.get("equipo_prueba") or request.data.get("equipo_configurado")
        metodo_equipo = request.data.get("metodo_equipo") or request.data.get("metodo_configurado")

        if not isinstance(values, list):
            return Response({"detail": "El campo values debe ser una lista."}, status=status.HTTP_400_BAD_REQUEST)

        fields = _build_result_fields(prueba_muestra)
        allowed_keys = {field["key"] for field in fields}
        field_by_key = {field["key"]: field for field in fields}

        clean_values = []
        errors = []

        for item in values:
            key = item.get("key")
            value = item.get("value")

            if key not in allowed_keys:
                errors.append({"key": key, "message": "El campo no pertenece a la estructura de la prueba."})
                continue

            if value in [None, ""] and estatus == "preliminar":
                errors.append({"key": key, "message": "El valor es obligatorio para confirmar."})
                continue

            message = _validate_result_value(field_by_key.get(key, {}), value)
            if message:
                errors.append({"key": key, "message": message})
                continue

            clean_values.append({
                "key": key,
                "value": value,
                "value_label": _display_value_for_field(field_by_key.get(key, {}), value),
                "observacion": item.get("observacion") or "",
            })

        if errors:
            return Response({"detail": "Hay errores de validación.", "errors": errors}, status=status.HTTP_400_BAD_REQUEST)

        if equipo_prueba:
            prueba_muestra.equipo_configurado = get_object_or_404(EquipoPrueba.objects.filter(activo=True), pk=equipo_prueba)
        if metodo_equipo:
            metodo = get_object_or_404(MetodoEquipo.objects.filter(activo=True), pk=metodo_equipo)
            if prueba_muestra.equipo_configurado_id and metodo.equipo_prueba_id != prueba_muestra.equipo_configurado_id:
                return Response({"detail": "El metodo no pertenece al equipo seleccionado."}, status=status.HTTP_400_BAD_REQUEST)
            prueba_muestra.metodo_configurado = metodo
            if not prueba_muestra.equipo_configurado_id:
                prueba_muestra.equipo_configurado = metodo.equipo_prueba

        payload = {
            "schema_version": 1,
            "mode": "dynamic_result_entry",
            "fields": fields,
            "values": clean_values,
            "saved_at": timezone.now().isoformat(),
            "saved_by": request.user.pk,
        }

        result = _active_result_for(prueba_muestra)
        fecha_dt = timezone.now()
        if fecha_medicion:
            try:
                fecha_dt = timezone.datetime.fromisoformat(str(fecha_medicion).replace("Z", "+00:00"))
            except Exception:
                fecha_dt = timezone.now()

        if result:
            if result.resultado != json.dumps(payload, ensure_ascii=False):
                HistoricoResultado.objects.create(
                    resultado=result,
                    resultado_anterior=result.resultado,
                    fecha_medicion_anterior=result.fecha_medicion,
                    usuario_medicion=result.usuario_medicion,
                    observaciones_anterior=result.observaciones,
                    usuario_modificacion=request.user,
                    motivo_cambio=request.data.get("motivo_cambio") or "Actualización de resultado",
                )
            result.resultado = json.dumps(payload, ensure_ascii=False)
            result.fecha_medicion = fecha_dt
            result.usuario_medicion = request.user
            result.estatus = estatus
            result.observaciones = observaciones or None
            result.save()
        else:
            result = Resultado.objects.create(
                prueba_muestra=prueba_muestra,
                resultado=json.dumps(payload, ensure_ascii=False),
                fecha_medicion=fecha_dt,
                usuario_medicion=request.user,
                estatus=estatus,
                observaciones=observaciones or None,
            )

        if estatus == "pendiente":
            prueba_muestra.estatus = "en_proceso"
            prueba_muestra.completada = False
        else:
            prueba_muestra.estatus = "completado"
            prueba_muestra.completada = True
            prueba_muestra.fecha_medicion = fecha_dt
            prueba_muestra.usuario_medicion = request.user

        prueba_muestra.valor = _compact_result_value(clean_values)
        prueba_muestra.observaciones = observaciones or prueba_muestra.observaciones
        prueba_muestra.save()
        _sync_sample_test_limit_evaluation(prueba_muestra, result)
        if not defer_rollups:
            _sync_sample_completion_from_tests(prueba_muestra.muestra)

            lote = prueba_muestra.muestra.lote
            if lote:
                lote.recalcular_estado(save=True)

        if compact_response:
            return Response({"detail": "Resultado guardado."})

        return Response({
            "detail": "Resultado guardado.",
            "resultado": ResultadoSerializer(result, context=self.get_serializer_context()).data,
            "prueba_muestra": PruebaMuestraSerializer(prueba_muestra, context=self.get_serializer_context()).data,
        })

    @action(detail=True, methods=['post'])
    def revisar(self, request, pk=None):
        resultado = self.get_object()
        estatus_nuevo = request.data.get('estatus')
        observaciones = request.data.get('observaciones', '')

        if not estatus_nuevo:
            return Response({'error': 'Se requiere el nuevo estatus'}, status=status.HTTP_400_BAD_REQUEST)

        if 'resultado' in request.data:
            HistoricoResultado.objects.create(
                resultado=resultado,
                resultado_anterior=resultado.resultado,
                fecha_medicion_anterior=resultado.fecha_medicion,
                usuario_medicion=resultado.usuario_medicion,
                observaciones_anterior=resultado.observaciones,
                usuario_modificacion=request.user,
                motivo_cambio=observaciones or 'Modificación durante revisión'
            )

        RevisionResultado.objects.create(
            resultado=resultado,
            usuario_revision=request.user,
            estatus_anterior=resultado.estatus,
            estatus_nuevo=estatus_nuevo,
            observaciones=observaciones
        )

        serializer = CreateResultadoSerializer(resultado, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data)

    @action(detail=False, methods=["get"])
    def entry_batches(self, request):
        search = request.query_params.get("search", "").strip()
        queryset = (
            self._scope_lotes(LoteMuestras.objects)
            .select_related("cliente_empresa", "tipo_gestion")
            .filter(muestras__resultados__estado_asignacion="confirmada")
            .distinct()
            .order_by("-fecha_actualizacion")
        )

        if search:
            queryset = queryset.filter(
                Q(id__icontains=search)
                | Q(cliente_empresa__nombre__icontains=search)
                | Q(cliente_ocasional_nombre__icontains=search)
                | Q(contacto_nombre__icontains=search)
            )

        page = self.paginate_queryset(queryset)
        if page is not None:
            return self.get_paginated_response([_serialize_entry_batch(lote) for lote in page])
        return Response([_serialize_entry_batch(lote) for lote in queryset])

    @action(detail=False, methods=["get"])
    def entry_batch_detail(self, request, lote_id=None):
        lote_id = lote_id or request.query_params.get("lote")
        lote = get_object_or_404(
            self._scope_lotes(LoteMuestras.objects.select_related("cliente_empresa", "tipo_gestion").prefetch_related(
                "muestras",
                "muestras__resultados",
                "muestras__resultados__prueba",
                "muestras__resultados__prueba__resultados",
            )),
            pk=lote_id,
        )

        sample_tests = (
            self._scope_sample_tests(PruebaMuestra.objects)
            .filter(muestra__lote=lote, estado_asignacion="confirmada")
            .select_related(
                "muestra",
                "prueba",
                "prueba__metodo",
                "prueba__metodo__equipo_prueba",
                "equipo_configurado",
                "metodo_configurado",
                "usuario_solicitud",
                "usuario_medicion",
            )
            .prefetch_related(
                "prueba__resultados__divisiones__componentes",
                "resultados",
            )
            .order_by("muestra_id", "prueba__nombre_variable")
        )

        data = _serialize_entry_batch(lote)
        data["sample_tests"] = PruebaMuestraSerializer(sample_tests, many=True, context=self.get_serializer_context()).data
        data["muestras"] = [
            {
                "id": muestra.id,
                "tipo_muestra": muestra.tipo_muestra,
                "condicion": muestra.condicion,
                "referencia_marca": muestra.referencia_marca,
                "equipo_placa": muestra.equipo_placa,
                "fecha_toma": muestra.fecha_toma,
                "is_ingresado": muestra.is_ingresado,
            }
            for muestra in lote.muestras.all()
        ]
        return Response(data)

    @action(detail=False, methods=["get"])
    def entry_form(self, request, prueba_muestra_id=None):
        prueba_muestra_id = prueba_muestra_id or request.query_params.get("prueba_muestra")
        prueba_muestra = get_object_or_404(
            self._scope_sample_tests(PruebaMuestra.objects.select_related(
                "muestra",
                "muestra__lote",
                "prueba",
                "prueba__metodo",
                "prueba__metodo__equipo_prueba",
                "equipo_configurado",
                "metodo_configurado",
            ).prefetch_related("prueba__resultados__divisiones__componentes")),
            pk=prueba_muestra_id,
        )
        result = _active_result_for(prueba_muestra)
        payload = _safe_json_loads(result.resultado) if result else None
        return Response({
            "prueba_muestra": PruebaMuestraSerializer(prueba_muestra, context=self.get_serializer_context()).data,
            "fields": _build_result_fields(prueba_muestra),
            "current_result": {
                "id": result.id,
                "estatus": result.estatus,
                "observaciones": result.observaciones,
                "fecha_medicion": result.fecha_medicion,
                "payload": payload,
            } if result else None,
        })

    @action(detail=False, methods=["post"])
    def save_draft(self, request, prueba_muestra_id=None):
        return self._save_result_payload(request, prueba_muestra_id, "pendiente")

    @action(detail=False, methods=["post"])
    def confirm_result(self, request, prueba_muestra_id=None):
        return self._save_result_payload(request, prueba_muestra_id, "preliminar")

    @action(detail=False, methods=["post"])
    def confirm_batch_results(self, request, lote_id=None):
        lote = get_object_or_404(self._scope_lotes(LoteMuestras.objects.all()), pk=lote_id)
        require_batch_operation(lote, "ingresar_resultados")
        sample_tests = list(
            self._scope_sample_tests(PruebaMuestra.objects)
            .filter(muestra__lote=lote, estado_asignacion="confirmada")
            .select_related("muestra", "prueba")
            .prefetch_related("prueba__resultados__divisiones__componentes")
            .order_by("muestra_id", "id")
        )
        if not sample_tests:
            return Response(
                {"detail": "El lote no tiene pruebas confirmadas para completar."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        results = {}
        incomplete = []
        for sample_test in sample_tests:
            result = _active_result_for(sample_test)
            payload = _safe_json_loads(result.resultado) if result else {}
            stored_values = payload.get("values", []) if isinstance(payload, dict) else []
            value_by_key = {
                item.get("key"): item.get("value")
                for item in stored_values
                if isinstance(item, dict) and item.get("key")
            }
            missing_fields = [
                field.get("label") or field.get("key")
                for field in _build_result_fields(sample_test)
                if value_by_key.get(field.get("key")) in [None, ""]
            ]
            if not result or missing_fields:
                incomplete.append({
                    "prueba_muestra": sample_test.id,
                    "muestra": str(sample_test.muestra_id),
                    "prueba": sample_test.prueba.nombre_variable,
                    "campos_faltantes": missing_fields or ["Resultado"],
                })
                continue
            results[sample_test.id] = result

        if incomplete:
            return Response({
                "detail": "Hay pruebas incompletas. Complete los campos indicados antes de continuar a revisión.",
                "incomplete": incomplete,
            }, status=status.HTTP_400_BAD_REQUEST)

        now = timezone.now()
        with transaction.atomic():
            Resultado.objects.filter(pk__in=[result.pk for result in results.values()]).update(
                estatus="preliminar",
                fecha_medicion=now,
                usuario_medicion=request.user,
            )
            PruebaMuestra.objects.filter(pk__in=results.keys()).update(
                estatus="completado",
                completada=True,
                fecha_medicion=now,
                usuario_medicion=request.user,
            )
            affected_samples = {sample_test.muestra_id: sample_test.muestra for sample_test in sample_tests}
            for sample in affected_samples.values():
                _sync_sample_completion_from_tests(sample)
            lote.recalcular_estado(save=True)

        return Response({
            "detail": "Todos los resultados fueron confirmados y están listos para revisión.",
            "tests_completed": len(results),
        })

    @action(detail=False, methods=["get"])
    def excel_template(self, request):
        lote_id = request.query_params.get("lote")
        mode = request.query_params.get("modo", "consolidado")
        if mode not in {"consolidado", "por_muestra"}:
            return Response(
                {"detail": "El modo debe ser consolidado o por_muestra."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        lote = get_object_or_404(self._scope_lotes(LoteMuestras.objects.all()), pk=lote_id)
        sample_tests = list(
            self._scope_sample_tests(PruebaMuestra.objects)
            .filter(muestra__lote=lote, estado_asignacion="confirmada")
            .select_related("muestra", "prueba")
            .prefetch_related("prueba__resultados__divisiones__componentes")
            .order_by("muestra_id", "prueba__nombre_variable")
        )

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill, Protection
            from openpyxl.utils import get_column_letter
            from openpyxl.workbook.defined_name import DefinedName
            from openpyxl.worksheet.datavalidation import DataValidation
        except Exception:
            return Response(
                {"detail": "openpyxl no está instalado en el entorno del backend."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        token = uuid.uuid4().hex
        wb = Workbook()
        wb.remove(wb.active)
        headers = [
            "Token",
            "Lote",
            "Muestra",
            "PruebaMuestraID",
            "Prueba",
            "Acronimo",
            "CampoKey",
            "Resultado",
            "Division",
            "Componente",
            "Unidad",
            "Valor",
            "Observaciones",
            "TipoCampo",
            "OpcionesPermitidas",
        ]

        lists = wb.create_sheet("_LISTAS")
        lists.sheet_state = "hidden"
        validation_counter = 0

        def safe_sheet_title(value, used):
            base = re.sub(r"[\\/*?:\[\]]", "_", str(value or "Muestra"))[:31] or "Muestra"
            title = base
            suffix = 2
            while title in used:
                marker = f"_{suffix}"
                title = f"{base[:31 - len(marker)]}{marker}"
                suffix += 1
            used.add(title)
            return title

        def add_option_validation(ws, cell, options):
            nonlocal validation_counter
            values = []
            for option in options or []:
                value = str(option.get("value", option)).strip()
                if value and value not in values:
                    values.append(value)
            if not values:
                return
            validation_counter += 1
            column = validation_counter
            for row_number, value in enumerate(values, start=1):
                lists.cell(row_number, column, value)
            name = f"opciones_{validation_counter}"
            reference = f"'_LISTAS'!${get_column_letter(column)}$1:${get_column_letter(column)}${len(values)}"
            defined_name = DefinedName(name, attr_text=reference)
            try:
                wb.defined_names.add(defined_name)
            except AttributeError:
                wb.defined_names.append(defined_name)
            validation = DataValidation(type="list", formula1=name, allow_blank=True)
            validation.error = "Seleccione un valor de la lista."
            validation.errorTitle = "Valor no permitido"
            validation.prompt = "Use una de las opciones configuradas para este campo."
            validation.promptTitle = "Campo controlado"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            ws.add_data_validation(validation)
            validation.add(cell)

        def configure_sheet(ws, tests):
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="222222")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.protection = Protection(locked=True)

            for st in tests:
                for field in _build_result_fields(st):
                    options = field.get("opciones") or []
                    ws.append([
                        token,
                        lote.id,
                        st.muestra_id,
                        st.id,
                        st.prueba.nombre_variable,
                        st.prueba.acronimo,
                        field["key"],
                        field["resultado_nombre"],
                        field["division_nombre"],
                        field["componente_nombre"],
                        field["unidad"],
                        "",
                        "",
                        field.get("tipo_dato") or "decimal",
                        " | ".join(str(option.get("label", option)) for option in options),
                    ])
                    row_number = ws.max_row
                    value_cell = ws.cell(row_number, 12)
                    observation_cell = ws.cell(row_number, 13)
                    value_cell.protection = Protection(locked=field.get("requiere_valor") is False)
                    observation_cell.protection = Protection(locked=field.get("permite_observacion") is False)

                    if options:
                        add_option_validation(ws, value_cell, options)
                    elif field.get("tipo_dato") in ["numerico", "decimal"]:
                        validation = DataValidation(
                            type="decimal",
                            operator="between",
                            formula1="-1E+100",
                            formula2="1E+100",
                            allow_blank=True,
                        )
                        validation.error = "Ingrese un valor numérico válido."
                        validation.errorTitle = "Formato incorrecto"
                        validation.showErrorMessage = True
                        ws.add_data_validation(validation)
                        validation.add(value_cell)
                    elif field.get("tipo_dato") in ["entero", "integer"]:
                        validation = DataValidation(
                            type="whole",
                            operator="between",
                            formula1="-1000000000000",
                            formula2="1000000000000",
                            allow_blank=True,
                        )
                        validation.error = "Ingrese un número entero válido."
                        validation.errorTitle = "Formato incorrecto"
                        validation.showErrorMessage = True
                        ws.add_data_validation(validation)
                        validation.add(value_cell)

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:O{max(ws.max_row, 1)}"
            ws.row_dimensions[1].height = 28
            widths = [34, 16, 18, 18, 30, 16, 48, 24, 24, 24, 14, 20, 34, 14, 36]
            for index, width in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(index)].width = width
            for hidden_column in [1, 4, 7, 14, 15]:
                ws.column_dimensions[get_column_letter(hidden_column)].hidden = True
            ws.protection.sheet = True
            ws.protection.password = "globaloil"
            ws.protection.autoFilter = False
            ws.protection.sort = False

        if mode == "consolidado":
            configure_sheet(wb.create_sheet("Resultados"), sample_tests)
        else:
            used_titles = {"Metadata", "_LISTAS"}
            tests_by_sample = defaultdict(list)
            for sample_test in sample_tests:
                tests_by_sample[str(sample_test.muestra_id)].append(sample_test)
            for sample_id, tests in tests_by_sample.items():
                configure_sheet(wb.create_sheet(safe_sheet_title(sample_id, used_titles)), tests)

        if not sample_tests:
            configure_sheet(wb.create_sheet("Resultados"), [])

        meta = wb.create_sheet("Metadata")
        meta.append(["token", token])
        meta.append(["lote", lote.id])
        meta.append(["modo", mode])
        meta.append(["generado_por", request.user.pk])
        meta.append(["fecha_generacion", timezone.now().isoformat()])
        meta.sheet_state = "hidden"
        for sheet_index, sheet in enumerate(wb.worksheets):
            if sheet.sheet_state == "visible":
                wb.active = sheet_index
                break

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)

        response = HttpResponse(
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        suffix = "una_hoja" if mode == "consolidado" else "hojas_por_muestra"
        response["Content-Disposition"] = f'attachment; filename="plantilla_resultados_{lote.id}_{suffix}.xlsx"'
        return response

    @action(detail=False, methods=["get"])
    def excel_export(self, request):
        lote_id = request.query_params.get("lote")
        mode = request.query_params.get("modo", "consolidado")
        sample_id = request.query_params.get("muestra")
        if mode not in {"consolidado", "por_muestra", "por_muestra_columnas"}:
            return Response(
                {"detail": "El modo debe ser consolidado, por_muestra o por_muestra_columnas."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        lote = get_object_or_404(self._scope_lotes(LoteMuestras.objects.all()), pk=lote_id)
        sample_tests = (
            self._scope_sample_tests(PruebaMuestra.objects)
            .filter(muestra__lote=lote, estado_asignacion="confirmada")
            .select_related("muestra", "muestra__referencia_equipo", "prueba")
            .prefetch_related("prueba__resultados__divisiones__componentes")
            .order_by("muestra_id", "prueba__nombre_variable", "id")
        )
        if sample_id:
            sample_tests = sample_tests.filter(muestra_id=sample_id)

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except Exception:
            return Response(
                {"detail": "openpyxl no esta instalado en el entorno del backend."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        tests = list(sample_tests)
        if not tests:
            return Response(
                {"detail": "No hay pruebas confirmadas para exportar."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        thin = Side(style="thin", color="B7B7B7")
        header_fill = PatternFill("solid", fgColor="242424")
        section_fill = PatternFill("solid", fgColor="E7E6E6")

        def style_sheet(ws, freeze="A2"):
            ws.freeze_panes = freeze
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column in range(1, ws.max_column + 1):
                values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
                width = min(max(max((len(value) for value in values), default=8) + 2, 12), 42)
                ws.column_dimensions[get_column_letter(column)].width = width

        def stored_values(st):
            result = _active_result_for(st)
            payload = _safe_json_loads(result.resultado) if result else None
            flattened = _flatten_payload_values(payload)
            by_key = {item.get("key"): item for item in flattened if item.get("key")}
            return result, flattened, by_key

        def field_value(field, flattened, by_key, index):
            item = by_key.get(field.get("key"))
            if item is None and index < len(flattened):
                item = flattened[index]
            return item or {}

        def excel_datetime(value):
            if not value:
                return ""
            if timezone.is_aware(value):
                value = timezone.localtime(value).replace(tzinfo=None)
            return value

        def evaluation_detail(field, details):
            result_id = field.get("resultado_id")
            component_id = field.get("componente_id")
            for detail in details:
                if (
                    str(detail.get("result_id") or "") == str(result_id or "")
                    and str(detail.get("component_id") or "") == str(component_id or "")
                ):
                    return detail
            field_key = str(field.get("codigo") or field.get("key") or "")
            return next((detail for detail in details if str(detail.get("field") or "") == field_key), {})

        wb = Workbook()
        wb.remove(wb.active)

        if mode == "consolidado":
            ws = wb.create_sheet("Resultados")
            samples = []
            tests_by_sample = {}
            for st in tests:
                tests_by_sample.setdefault(st.muestra_id, []).append(st)
                if st.muestra_id not in samples:
                    samples.append(st.muestra_id)

            column_specs = []
            seen_specs = set()
            for st in tests:
                for field in _build_result_fields(st):
                    spec_key = (st.prueba_id, field.get("key"))
                    if spec_key in seen_specs:
                        continue
                    seen_specs.add(spec_key)
                    result_name = field.get("resultado_nombre") or "Resultado"
                    field_name = field.get("componente_nombre") or field.get("label") or "Valor"
                    column_specs.append((spec_key, f"{st.prueba.acronimo} | {result_name} | {field_name}"))

            headers = ["Lote", "Muestra", "Fecha de toma", "Equipo", "Referencia / marca"] + [label for _, label in column_specs]
            ws.append(headers)
            for current_sample_id in samples:
                sample_group = tests_by_sample[current_sample_id]
                sample = sample_group[0].muestra
                values_by_spec = {}
                for st in sample_group:
                    _, flattened, by_key = stored_values(st)
                    for index, field in enumerate(_build_result_fields(st)):
                        item = field_value(field, flattened, by_key, index)
                        display = item.get("value_label") or item.get("value") or ""
                        values_by_spec[(st.prueba_id, field.get("key"))] = display
                ws.append([
                    lote.id,
                    sample.id,
                    excel_datetime(sample.fecha_toma),
                    _sample_equipment_label(sample),
                    getattr(sample, "referencia_marca", "") or getattr(sample, "descripcion", "") or "",
                    *[values_by_spec.get(spec_key, "") for spec_key, _ in column_specs],
                ])
            style_sheet(ws)
        elif mode == "por_muestra":
            tests_by_sample = {}
            for st in tests:
                tests_by_sample.setdefault(st.muestra_id, []).append(st)
            for sample_id_key, sample_group in tests_by_sample.items():
                sample = sample_group[0].muestra
                raw_title = str(sample.id or sample_id_key)
                title = "".join("_" if char in '[]:*?/\\' else char for char in raw_title)[:31] or "Muestra"
                base_title = title
                suffix = 2
                while title in wb.sheetnames:
                    title = f"{base_title[:27]}_{suffix}"
                    suffix += 1
                ws = wb.create_sheet(title)
                ws.append(["Prueba", "Resultado", "Campo", "Estructura", "Valor", "Unidad", "Observacion"])
                for st in sample_group:
                    _, flattened, by_key = stored_values(st)
                    for index, field in enumerate(_build_result_fields(st)):
                        item = field_value(field, flattened, by_key, index)
                        result_name = field.get("resultado_nombre") or "Resultado"
                        component_name = field.get("componente_nombre") or field.get("label") or "Valor"
                        structure = " / ".join(filter(None, [field.get("division_nombre"), component_name]))
                        ws.append([
                            f"{st.prueba.acronimo} - {st.prueba.nombre_variable}",
                            result_name,
                            component_name,
                            structure,
                            item.get("value_label") or item.get("value") or "",
                            field.get("unidad") or item.get("unidad") or "",
                            item.get("observacion") or "",
                        ])
                style_sheet(ws)
                ws["A1"].fill = header_fill
                for row in range(2, ws.max_row + 1):
                    if row == 2 or ws.cell(row, 1).value != ws.cell(row - 1, 1).value:
                        for cell in ws[row]:
                            cell.fill = section_fill
        else:
            tests_by_sample = {}
            for st in tests:
                tests_by_sample.setdefault(st.muestra_id, []).append(st)

            for sample_id_key, sample_group in tests_by_sample.items():
                sample = sample_group[0].muestra
                raw_title = str(sample.id or sample_id_key)
                title = "".join("_" if char in '[]:*?/\\' else char for char in raw_title)[:31] or "Muestra"
                base_title = title
                title_suffix = 2
                while title in wb.sheetnames:
                    title = f"{base_title[:27]}_{title_suffix}"
                    title_suffix += 1

                ws = wb.create_sheet(title)
                specs = []
                row_values = {}
                for st in sample_group:
                    result, flattened, by_key = stored_values(st)
                    evaluation = _evaluate_sample_test_limits(st, result)
                    details = evaluation.get("details") or []
                    for index, field in enumerate(_build_result_fields(st)):
                        item = field_value(field, flattened, by_key, index)
                        detail = evaluation_detail(field, details)
                        result_name = field.get("resultado_nombre") or "Resultado"
                        field_name = field.get("componente_nombre") or field.get("label") or "Valor"
                        base_label = f"{st.prueba.acronimo} | {result_name} | {field_name}"
                        specs.extend([
                            (f"value:{st.id}:{index}", base_label),
                            (f"status:{st.id}:{index}", f"Estado | {base_label}"),
                            (f"rule:{st.id}:{index}", f"Regla | {base_label}"),
                        ])
                        row_values[f"value:{st.id}:{index}"] = item.get("value_label") or item.get("value") or ""
                        row_values[f"status:{st.id}:{index}"] = _evaluation_status_label(
                            _canonical_status(detail.get("estado"))
                        ) if detail else "NO EVALUABLE"
                        row_values[f"rule:{st.id}:{index}"] = _limit_detail_label(detail) if detail else "Sin regla aplicable"

                headers = [
                    "Lote", "Muestra", "Fecha de toma", "Equipo", "Referencia / marca", "Estado de la muestra",
                    *[label for _key, label in specs],
                ]
                sample_review_tests = [
                    {
                        "estado_tecnico": _canonical_status(
                            _evaluate_sample_test_limits(st, _active_result_for(st)).get("estado")
                        )
                    }
                    for st in sample_group
                ]
                ws.append(headers)
                ws.append([
                    lote.id,
                    sample.id,
                    excel_datetime(sample.fecha_toma),
                    _sample_equipment_label(sample),
                    getattr(sample, "referencia_marca", "") or getattr(sample, "descripcion", "") or "",
                    _sample_review_status(sample, sample_review_tests),
                    *[row_values.get(key, "") for key, _label in specs],
                ])
                style_sheet(ws)

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        suffix = {
            "consolidado": "consolidado",
            "por_muestra": "por_muestra",
            "por_muestra_columnas": "por_muestra_columnas",
        }[mode]
        response = HttpResponse(
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="resultados_{lote.id}_{suffix}.xlsx"'
        return response

    def _legacy_excel_preview(self, request):
        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return Response({"detail": "Debe adjuntar un archivo Excel en el campo file."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            from openpyxl import load_workbook
        except Exception:
            return Response({"detail": "openpyxl no está instalado en el entorno del backend."}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        try:
            wb = load_workbook(uploaded_file, data_only=True)
            ws = wb["Resultados"]
        except Exception:
            return Response({"detail": "El archivo no contiene una hoja Resultados válida."}, status=status.HTTP_400_BAD_REQUEST)

        required = [
            "Token", "Lote", "Muestra", "PruebaMuestraID", "CampoKey", "Valor"
        ]
        headers = [cell.value for cell in ws[1]]
        missing = [name for name in required if name not in headers]

        if missing:
            return Response({"detail": "La plantilla no tiene las columnas requeridas.", "missing": missing}, status=status.HTTP_400_BAD_REQUEST)

        index = {name: headers.index(name) + 1 for name in headers if name}
        errors = []
        valid_rows = []

        for row_number in range(2, ws.max_row + 1):
            prueba_muestra_id = ws.cell(row_number, index["PruebaMuestraID"]).value
            field_key = ws.cell(row_number, index["CampoKey"]).value
            value = ws.cell(row_number, index["Valor"]).value
            lote_id = ws.cell(row_number, index["Lote"]).value

            if value in [None, ""]:
                continue

            try:
                prueba_muestra = self._scope_sample_tests(
                    PruebaMuestra.objects.select_related("muestra", "muestra__lote", "prueba")
                ).get(pk=prueba_muestra_id)
            except PruebaMuestra.DoesNotExist:
                errors.append({"row": row_number, "message": "La prueba asignada no existe."})
                continue

            if str(prueba_muestra.muestra.lote_id) != str(lote_id):
                errors.append({"row": row_number, "message": "La muestra no pertenece al lote indicado."})
                continue

            allowed = {field["key"] for field in _build_result_fields(prueba_muestra)}
            if field_key not in allowed:
                errors.append({"row": row_number, "message": "El campo no pertenece a la estructura de la prueba."})
                continue

            valid_rows.append({
                "row": row_number,
                "prueba_muestra": prueba_muestra_id,
                "field_key": field_key,
                "value": value,
            })

        return Response({
            "valid": not errors,
            "rows_valid": len(valid_rows),
            "rows_error": len(errors),
            "errors": errors[:200],
            "preview": valid_rows[:50],
        })

    def _parse_result_excel(self, uploaded_file):
        try:
            from openpyxl import load_workbook
        except Exception:
            return [], [], Response(
                {"detail": "openpyxl no está instalado en el entorno del backend."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        try:
            wb = load_workbook(uploaded_file, data_only=True)
        except Exception:
            return [], [], Response(
                {"detail": "No se pudo leer el archivo Excel. Use una plantilla oficial .xlsx."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        ignored_sheets = {"Metadata", "_LISTAS", "Instrucciones"}
        data_sheets = [ws for ws in wb.worksheets if ws.title not in ignored_sheets and ws.sheet_state == "visible"]
        if not data_sheets:
            return [], [], Response(
                {"detail": "El archivo no contiene hojas de resultados válidas."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        metadata_token = None
        if "Metadata" in wb.sheetnames:
            metadata = {
                str(wb["Metadata"].cell(row, 1).value): wb["Metadata"].cell(row, 2).value
                for row in range(1, wb["Metadata"].max_row + 1)
            }
            metadata_token = metadata.get("token")
        if not metadata_token:
            return [], [], Response(
                {"detail": "El archivo no conserva los metadatos de una plantilla oficial."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        required = ["Token", "Lote", "Muestra", "PruebaMuestraID", "CampoKey", "Valor"]
        errors = []
        valid_rows = []
        test_cache = {}
        field_cache = {}
        seen_fields = set()

        for ws in data_sheets:
            headers = [cell.value for cell in ws[1]]
            missing = [name for name in required if name not in headers]
            if missing:
                errors.append({
                    "sheet": ws.title,
                    "row": 1,
                    "message": f"Faltan columnas requeridas: {', '.join(missing)}.",
                })
                continue

            index = {name: headers.index(name) + 1 for name in headers if name}
            for row_number in range(2, ws.max_row + 1):
                prueba_muestra_id = ws.cell(row_number, index["PruebaMuestraID"]).value
                field_key = ws.cell(row_number, index["CampoKey"]).value
                value = ws.cell(row_number, index["Valor"]).value
                lote_id = ws.cell(row_number, index["Lote"]).value
                sample_id = ws.cell(row_number, index["Muestra"]).value
                row_token = ws.cell(row_number, index["Token"]).value
                observation = ws.cell(row_number, index.get("Observaciones", 0)).value if index.get("Observaciones") else ""

                if value in [None, ""]:
                    continue
                if metadata_token and str(row_token) != str(metadata_token):
                    errors.append({"sheet": ws.title, "row": row_number, "message": "El token de la fila fue alterado."})
                    continue

                if prueba_muestra_id not in test_cache:
                    try:
                        test_cache[prueba_muestra_id] = self._scope_sample_tests(
                            PruebaMuestra.objects.select_related("muestra", "muestra__lote", "prueba")
                            .prefetch_related("prueba__resultados__divisiones__componentes")
                        ).get(pk=prueba_muestra_id, estado_asignacion="confirmada")
                    except (PruebaMuestra.DoesNotExist, TypeError, ValueError):
                        test_cache[prueba_muestra_id] = None

                prueba_muestra = test_cache.get(prueba_muestra_id)
                if not prueba_muestra:
                    errors.append({"sheet": ws.title, "row": row_number, "message": "La prueba asignada no existe o no está confirmada."})
                    continue
                if str(prueba_muestra.muestra.lote_id) != str(lote_id):
                    errors.append({"sheet": ws.title, "row": row_number, "message": "La muestra no pertenece al lote indicado."})
                    continue
                if str(prueba_muestra.muestra_id) != str(sample_id):
                    errors.append({"sheet": ws.title, "row": row_number, "message": "La fila no corresponde a la muestra indicada."})
                    continue

                if prueba_muestra_id not in field_cache:
                    field_cache[prueba_muestra_id] = {field["key"]: field for field in _build_result_fields(prueba_muestra)}
                field = field_cache[prueba_muestra_id].get(field_key)
                if not field:
                    errors.append({"sheet": ws.title, "row": row_number, "message": "El campo no pertenece a la estructura actual de la prueba."})
                    continue

                validation_message = _validate_result_value(field, value)
                if validation_message:
                    errors.append({"sheet": ws.title, "row": row_number, "message": validation_message})
                    continue

                value = _normalize_imported_result_value(field, value)

                field_identity = (prueba_muestra_id, field_key)
                if field_identity in seen_fields:
                    errors.append({"sheet": ws.title, "row": row_number, "message": "El campo está repetido en el archivo."})
                    continue
                seen_fields.add(field_identity)

                valid_rows.append({
                    "sheet": ws.title,
                    "row": row_number,
                    "prueba_muestra": prueba_muestra_id,
                    "field_key": field_key,
                    "value": value,
                    "observacion": observation or "",
                })

        return valid_rows, errors, None

    @action(detail=False, methods=["post"])
    def excel_preview(self, request):
        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return Response({"detail": "Debe adjuntar un archivo Excel en el campo file."}, status=status.HTTP_400_BAD_REQUEST)
        valid_rows, errors, fatal_response = self._parse_result_excel(uploaded_file)
        if fatal_response:
            return fatal_response
        values_by_test = defaultdict(list)
        for row in valid_rows:
            values_by_test[str(row["prueba_muestra"])].append({
                "key": row["field_key"],
                "value": row["value"],
                "observacion": row["observacion"],
            })
        return Response({
            "valid": not errors,
            "rows_valid": len(valid_rows),
            "rows_error": len(errors),
            "tests_valid": len({row["prueba_muestra"] for row in valid_rows}),
            "errors": errors[:200],
            "preview": valid_rows[:50],
            "values_by_test": values_by_test,
        })

    @action(detail=False, methods=["post"])
    def excel_import(self, request):
        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return Response({"detail": "Debe adjuntar un archivo Excel en el campo file."}, status=status.HTTP_400_BAD_REQUEST)
        valid_rows, errors, fatal_response = self._parse_result_excel(uploaded_file)
        if fatal_response:
            return fatal_response
        if errors:
            return Response({
                "detail": "El archivo contiene errores. Corríjalos antes de aplicar los resultados.",
                "rows_valid": len(valid_rows),
                "rows_error": len(errors),
                "errors": errors[:200],
            }, status=status.HTTP_400_BAD_REQUEST)
        if not valid_rows:
            return Response({"detail": "El archivo no contiene valores para importar."}, status=status.HTTP_400_BAD_REQUEST)

        grouped = defaultdict(dict)
        for row in valid_rows:
            grouped[row["prueba_muestra"]][row["field_key"]] = {
                "key": row["field_key"],
                "value": row["value"],
                "observacion": row["observacion"],
            }

        with transaction.atomic():
            affected_samples = {}
            affected_batches = {}
            for prueba_muestra_id, imported_values in grouped.items():
                prueba_muestra = PruebaMuestra.objects.select_related(
                    "muestra",
                    "muestra__lote",
                ).get(pk=prueba_muestra_id)
                affected_samples[prueba_muestra.muestra_id] = prueba_muestra.muestra
                if prueba_muestra.muestra.lote_id:
                    affected_batches[prueba_muestra.muestra.lote_id] = prueba_muestra.muestra.lote
                current_result = _active_result_for(prueba_muestra)
                current_payload = _safe_json_loads(current_result.resultado) if current_result else {}
                merged_values = {
                    item.get("key"): {
                        "key": item.get("key"),
                        "value": item.get("value"),
                        "observacion": item.get("observacion") or "",
                    }
                    for item in (current_payload.get("values", []) if isinstance(current_payload, dict) else [])
                    if item.get("key")
                }
                merged_values.update(imported_values)
                proxy_request = SimpleNamespace(data={
                    "values": list(merged_values.values()),
                    "observaciones": current_result.observaciones if current_result else "",
                    "motivo_cambio": "Carga masiva desde plantilla Excel",
                }, user=request.user)
                save_response = self._save_result_payload(
                    proxy_request,
                    prueba_muestra_id,
                    "pendiente",
                    defer_rollups=True,
                    compact_response=True,
                )
                if save_response.status_code >= 400:
                    transaction.set_rollback(True)
                    return save_response

            for sample in affected_samples.values():
                _sync_sample_completion_from_tests(sample)
            for batch in affected_batches.values():
                batch.recalcular_estado(save=True)

        return Response({
            "detail": "Resultados aplicados como borrador.",
            "rows_imported": len(valid_rows),
            "tests_updated": len(grouped),
        })

    @action(detail=False, methods=["get"])
    def review_batches(self, request):
        search = request.query_params.get("search", "").strip()
        queryset = (
            self._scope_lotes(
                LoteMuestras.objects
                .select_related("cliente_empresa", "tipo_gestion")
                .filter(muestras__resultados__resultados__isnull=False)
            )
            .distinct()
            .order_by("-fecha_actualizacion")
        )
        if search:
            queryset = queryset.filter(
                Q(id__icontains=search)
                | Q(cliente_empresa__nombre__icontains=search)
                | Q(cliente_ocasional_nombre__icontains=search)
                | Q(contacto_nombre__icontains=search)
            )

        page = self.paginate_queryset(queryset)
        if page is not None:
            return self.get_paginated_response([_serialize_entry_batch(lote) for lote in page])
        return Response([_serialize_entry_batch(lote) for lote in queryset])

    @action(detail=False, methods=["get"])
    def review_batch_detail(self, request, lote_id=None):
        lote_id = lote_id or request.query_params.get("lote")
        lote = get_object_or_404(
            LoteMuestras.objects.select_related("cliente_empresa", "tipo_gestion").prefetch_related("muestras"),
            pk=lote_id,
        )
        return Response(_batch_review_payload(lote))

    @action(detail=False, methods=["get"])
    def result_history(self, request, resultado_id=None):
        resultado = get_object_or_404(Resultado.objects.select_related("usuario_medicion", "prueba_muestra", "prueba_muestra__prueba"), pk=resultado_id)
        history = []
        history.append({
            "id": f"initial-{resultado.id}",
            "type": "initial",
            "fecha": resultado.fecha_registro,
            "usuario": _user_payload(resultado.usuario_medicion),
            "valor_anterior": None,
            "valor_nuevo": _result_summary(resultado),
            "justificacion": "Resultado inicial registrado.",
            "payload": _safe_json_loads(resultado.resultado),
        })
        for item in resultado.historico.select_related("usuario_modificacion", "usuario_medicion").order_by("-fecha_modificacion"):
            old_payload = _safe_json_loads(item.resultado_anterior)
            history.append({
                "id": item.id,
                "type": "correction",
                "fecha": item.fecha_modificacion,
                "usuario": _user_payload(item.usuario_modificacion),
                "valor_anterior": " / ".join([str(v.get("value_label") or v.get("value")) for v in _flatten_payload_values(old_payload) if v.get("value") not in [None, ""]]) or item.resultado_anterior,
                "valor_nuevo": _result_summary(resultado),
                "justificacion": item.motivo_cambio,
                "payload": old_payload,
            })
        revisions = [
            {
                "id": f"revision-{rev.id}",
                "type": "revision",
                "fecha": rev.fecha_revision,
                "usuario": _user_payload(rev.usuario_revision),
                "valor_anterior": rev.estatus_anterior,
                "valor_nuevo": rev.estatus_nuevo,
                "justificacion": rev.observaciones,
                "payload": None,
            }
            for rev in resultado.revisiones.select_related("usuario_revision").order_by("-fecha_revision")
        ]
        return Response({"resultado": ResultadoSerializer(resultado, context=self.get_serializer_context()).data, "history": history + revisions})

    @action(detail=False, methods=["post"])
    def correct_result(self, request, resultado_id=None):
        resultado = get_object_or_404(Resultado.objects.select_related("prueba_muestra", "prueba_muestra__muestra", "prueba_muestra__prueba"), pk=resultado_id)
        justificacion = (request.data.get("justificacion") or request.data.get("motivo_cambio") or "").strip()
        if not justificacion:
            return Response({"detail": "La justificación es obligatoria para corregir un resultado."}, status=status.HTTP_400_BAD_REQUEST)

        values = request.data.get("values", [])
        if not isinstance(values, list):
            return Response({"detail": "El campo values debe ser una lista."}, status=status.HTTP_400_BAD_REQUEST)

        old_payload = _safe_json_loads(resultado.resultado)
        fields = old_payload.get("fields") if isinstance(old_payload, dict) else _build_result_fields(resultado.prueba_muestra)
        allowed = {field.get("key") for field in fields if field.get("key")}
        field_by_key = {field.get("key"): field for field in fields if field.get("key")}
        clean_values = []
        for item in values:
            key = item.get("key")
            if key not in allowed:
                return Response({"detail": f"El campo {key} no pertenece a la estructura del resultado."}, status=status.HTTP_400_BAD_REQUEST)
            message = _validate_result_value(field_by_key.get(key, {}), item.get("value"))
            if message:
                return Response({"detail": message, "key": key}, status=status.HTTP_400_BAD_REQUEST)
            clean_values.append({
                "key": key,
                "value": item.get("value"),
                "value_label": _display_value_for_field(field_by_key.get(key, {}), item.get("value")),
                "observacion": item.get("observacion") or "",
            })

        HistoricoResultado.objects.create(
            resultado=resultado,
            resultado_anterior=resultado.resultado,
            fecha_medicion_anterior=resultado.fecha_medicion,
            usuario_medicion=resultado.usuario_medicion,
            observaciones_anterior=resultado.observaciones,
            usuario_modificacion=request.user,
            motivo_cambio=justificacion,
        )

        new_payload = {
            "schema_version": 1,
            "mode": "review_correction",
            "fields": fields,
            "values": clean_values,
            "corrected_at": timezone.now().isoformat(),
            "corrected_by": request.user.pk,
            "justificacion": justificacion,
        }
        resultado.resultado = json.dumps(new_payload, ensure_ascii=False)
        resultado.fecha_actualizacion = timezone.now()
        resultado.usuario_medicion = request.user
        resultado.observaciones = request.data.get("observaciones") or resultado.observaciones
        resultado.save()

        st = resultado.prueba_muestra
        st.valor = _compact_result_value(clean_values)
        st.usuario_medicion = request.user
        st.fecha_medicion = timezone.now()
        st.estatus = "completado"
        st.completada = True
        st.save()
        _sync_sample_test_limit_evaluation(st, resultado)
        _sync_sample_completion_from_tests(st.muestra)
        if st.muestra and st.muestra.lote:
            st.muestra.lote.recalcular_estado(save=True)

        return Response({"detail": "Resultado corregido.", "resultado": ResultadoSerializer(resultado, context=self.get_serializer_context()).data})

    @action(detail=False, methods=["post"])
    def mark_sample_test_reviewed(self, request, prueba_muestra_id=None):
        prueba_muestra = get_object_or_404(PruebaMuestra.objects.select_related("muestra", "muestra__lote"), pk=prueba_muestra_id)
        require_batch_operation(prueba_muestra.muestra.lote, "revisar_resultados")
        if not prueba_muestra.completada:
            return Response(
                {"detail": "No se puede revisar una prueba sin resultado completo."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        prueba_muestra.usuario_revision = request.user
        prueba_muestra.fecha_revision = timezone.now()
        prueba_muestra.estatus = "aprobado"
        prueba_muestra.is_revisada = True
        prueba_muestra.save()

        result = _active_result_for(prueba_muestra)
        if result:
            RevisionResultado.objects.create(
                resultado=result,
                usuario_revision=request.user,
                estatus_anterior=result.estatus,
                estatus_nuevo="aprobado",
                observaciones=request.data.get("observaciones") or "Resultado revisado.",
            )
            result.estatus = "aprobado"
            result.save(update_fields=["estatus", "fecha_actualizacion"])

        muestra = prueba_muestra.muestra
        _sync_sample_completion_from_tests(muestra)
        if muestra.is_revisado and muestra.lote:
            muestra.lote.recalcular_estado(save=True)

        return Response({"detail": "Resultado marcado como revisado."})

    @action(detail=False, methods=["post"])
    def complete_batch_review(self, request, lote_id=None):
        lote = get_object_or_404(LoteMuestras.objects.prefetch_related("muestras"), pk=lote_id)
        require_batch_operation(lote, "revisar_resultados")
        incomplete_count = PruebaMuestra.objects.filter(
            muestra__lote=lote,
            estado_asignacion="confirmada",
        ).exclude(completada=True).count()
        if incomplete_count:
            return Response(
                {"detail": f"No se puede revisar el lote: faltan {incomplete_count} prueba(s) por completar."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        tests = (
            PruebaMuestra.objects.filter(
                muestra__lote=lote,
                estado_asignacion="confirmada",
                resultados__isnull=False,
            )
            .select_related("muestra")
            .prefetch_related("resultados")
            .distinct()
        )
        now = timezone.now()
        count = 0
        for st in tests:
            if not st.is_revisada:
                st.is_revisada = True
                st.usuario_revision = request.user
                st.fecha_revision = now
                st.estatus = "aprobado"
                st.save()
                result = _active_result_for(st)
                if result:
                    RevisionResultado.objects.create(
                        resultado=result,
                        usuario_revision=request.user,
                        estatus_anterior=result.estatus,
                        estatus_nuevo="aprobado",
                        observaciones=request.data.get("observaciones") or "Revisión masiva del lote.",
                    )
                count += 1

        for muestra in lote.muestras.all():
            _sync_sample_completion_from_tests(muestra)

        lote.recalcular_estado(save=True)
        return Response({"detail": "Revisión del lote completada.", "updated": count})

    @action(detail=False, methods=["get"])
    def interpretation_batch_detail(self, request, lote_id=None):
        lote_id = lote_id or request.query_params.get("lote")
        lote = get_object_or_404(
            LoteMuestras.objects.select_related("cliente_empresa", "tipo_gestion").prefetch_related("muestras"),
            pk=lote_id,
        )
        return Response(_batch_review_payload(lote))

    @action(detail=False, methods=["get"])
    def interpretation_sample(self, request, muestra_id=None):
        from apps.muestras.api.models.muestras.index import Muestra
        muestra = get_object_or_404(Muestra.objects.select_related("lote", "referencia_equipo"), pk=muestra_id)
        tests = _sample_tests_for_lote(
            muestra.lote,
            muestra.id,
            include_test_structure=False,
        ) if muestra.lote_id else []
        for test in tests:
            # La pantalla consume el resumen y los detalles normalizados. El JSON
            # crudo del resultado y la evaluación repetían la misma información
            # varias veces y hacían que cada cambio de muestra transfiriera cientos
            # de KB innecesarios.
            result = test.get("resultado")
            if isinstance(result, dict):
                result.pop("resultado", None)
                result.pop("payload", None)
            evaluation = test.get("evaluacion_limite")
            if isinstance(evaluation, dict):
                evaluation.pop("details", None)
                evaluation.pop("raw", None)
        interpretation = _get_interpretation_data(muestra)
        return Response({
            "muestra": {
                "id": muestra.id,
                "lote": muestra.lote_id,
                "cliente_nombre": getattr(muestra.lote.cliente_empresa, "nombre", None) or getattr(muestra.lote.cliente_empresa, "name", None) or muestra.lote.cliente_ocasional_nombre if muestra.lote_id else None,
                "fecha_toma": muestra.fecha_toma,
                "equipo": _sample_equipment_label(muestra),
                "tipo_muestra": muestra.tipo_muestra,
                "referencia_marca": muestra.referencia_marca,
                "is_revisado": muestra.is_revisado,
            },
            "sample_tests": tests,
            "predefined_comments": _build_predefined_comments(tests),
            "selected_predefined_comments": interpretation.get("comentarios_predefinidos") if "comentarios_predefinidos" in interpretation else _build_predefined_comments(tests),
            "conclusion": interpretation.get("conclusion") or interpretation.get("comentario_proveedor") or "",
            "comentarios_manuales": interpretation.get("comentarios_manuales") or [],
            "graficas_activas": interpretation.get("graficas_activas", True),
            "graficas_seleccionadas": interpretation.get("graficas_seleccionadas") or [t["prueba"]["acronimo"] for t in tests[:4]],
            "interpretacion": interpretation,
        })

    @action(detail=False, methods=["patch", "post"])
    def save_interpretation_sample(self, request, muestra_id=None):
        from apps.muestras.api.models.muestras.index import Muestra
        muestra = get_object_or_404(Muestra, pk=muestra_id)
        data = {
            "conclusion": request.data.get("conclusion") or "",
            "comentarios_predefinidos": request.data.get("comentarios_predefinidos") or [],
            "comentarios_manuales": request.data.get("comentarios_manuales") or [],
            "graficas_activas": bool(request.data.get("graficas_activas", True)),
            "graficas_seleccionadas": request.data.get("graficas_seleccionadas") or [],
            "updated_at": timezone.now().isoformat(),
            "updated_by": request.user.pk,
        }
        _set_interpretation_data(muestra, data)
        return Response({"detail": "Interpretación guardada.", "interpretacion": data})

    @action(detail=False, methods=["get"])
    def interpretation_trends(self, request, muestra_id=None):
        from apps.muestras.api.models.muestras.index import Muestra
        muestra = get_object_or_404(Muestra.objects.select_related("lote", "referencia_equipo"), pk=muestra_id)
        prueba_filter = request.query_params.getlist("pruebas") or []
        related_samples = Muestra.objects.filter(lote__cliente_empresa=muestra.lote.cliente_empresa) if muestra.lote_id else Muestra.objects.all()
        if muestra.referencia_equipo_id:
            related_samples = related_samples.filter(referencia_equipo_id=muestra.referencia_equipo_id)
        elif muestra.equipo_placa:
            related_samples = related_samples.filter(equipo_placa=muestra.equipo_placa)
        related_samples = related_samples.order_by("-fecha_toma", "-id")[:100]

        tests = (
            PruebaMuestra.objects
            .filter(muestra__in=related_samples, resultados__isnull=False)
            .select_related("muestra", "prueba")
            .prefetch_related("resultados")
            .distinct()
            .order_by("muestra__fecha_toma", "muestra_id", "prueba__nombre_variable", "prueba_id")
        )
        groups_map = {}
        for st in tests:
            acronym = st.prueba.acronimo
            if prueba_filter and acronym not in prueba_filter:
                continue
            result = _active_result_for(st)
            payload = _safe_json_loads(result.resultado) if result else None
            flat = _flatten_payload_values(payload)
            for item in flat:
                raw_value = item.get("value")
                if raw_value in [None, ""]:
                    continue

                result_name = item.get("resultado_nombre") or "Resultado principal"
                field_key = item.get("key") or item.get("label")
                field_name = item.get("componente_nombre") or item.get("label") or "Valor"
                test_group = groups_map.setdefault(acronym, {
                    "prueba_id": st.prueba_id,
                    "prueba": acronym,
                    "nombre": st.prueba.nombre_variable,
                    "resultados": {},
                })
                result_group = test_group["resultados"].setdefault(result_name, {
                    "resultado": result_name,
                    "campos": {},
                })
                field_group = result_group["campos"].setdefault(field_key, {
                    "key": field_key,
                    "campo": field_name,
                    "unidad": item.get("unidad") or "",
                    "tipo": "categorico",
                    "points": [],
                })

                try:
                    numeric_value = float(str(raw_value).replace(",", "."))
                    field_group["tipo"] = "numerico"
                    point_value = numeric_value
                except Exception:
                    point_value = item.get("value_label") or raw_value

                field_group["points"].append({
                    "fecha": st.muestra.fecha_toma,
                    "muestra": st.muestra_id,
                    "valor": point_value,
                    "etiqueta": item.get("value_label") or str(raw_value),
                    "actual": st.muestra_id == muestra.id,
                })

        groups = []
        legacy_numeric_series = []
        for test_group in groups_map.values():
            serialized_results = []
            for result_group in test_group["resultados"].values():
                numeric_series = []
                categorical_series = []
                for field_group in result_group["campos"].values():
                    target = numeric_series if field_group["tipo"] == "numerico" else categorical_series
                    target.append(field_group)
                    if field_group["tipo"] == "numerico":
                        legacy_numeric_series.append({
                            "prueba": test_group["prueba"],
                            "nombre": test_group["nombre"],
                            "resultado": result_group["resultado"],
                            "campo": field_group["campo"],
                            "unidad": field_group["unidad"],
                            "points": field_group["points"],
                        })
                serialized_results.append({
                    "resultado": result_group["resultado"],
                    "numeric_series": numeric_series,
                    "categorical_series": categorical_series,
                })
            groups.append({
                "prueba_id": test_group["prueba_id"],
                "prueba": test_group["prueba"],
                "nombre": test_group["nombre"],
                "resultados": serialized_results,
            })

        return Response({
            "muestra": muestra.id,
            "equipo": _sample_equipment_label(muestra),
            "groups": groups,
            "series": legacy_numeric_series,
        })


class HistoricoResultadoViewSet(viewsets.ReadOnlyModelViewSet):
    permission_classes = [IsAuthenticated]
    queryset = HistoricoResultado.objects.select_related(
        'resultado',
        'resultado__prueba_muestra',
        'resultado__prueba_muestra__prueba',
        'usuario_medicion',
        'usuario_modificacion',
    )
    serializer_class = HistoricoResultadoSerializer


class RevisionResultadoViewSet(viewsets.ReadOnlyModelViewSet):
    pagination_class = StandardResultsSetPagination
    permission_classes = [IsAuthenticated]
    queryset = RevisionResultado.objects.select_related(
        'resultado',
        'resultado__prueba_muestra',
        'resultado__prueba_muestra__prueba',
        'usuario_revision',
    )
    serializer_class = RevisionResultadoSerializer
