import re
import uuid
from datetime import timedelta
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from permissions import ActionPermissionMixin, DenyReadOnlyWrite, HasSecurityPermission

from apps.activesTree.api.models.machines.index import Maquina
from apps.misc.api.models.dynamicTechnicalConfig.index import CampoTecnicoMuestra
from apps.muestras.api.models.sampleBatchExcel.index import SampleBatchExcelTemplateToken
from apps.users.api.models.index import User

SCHEMA_VERSION = "sample-batch-v4-dynamic-technical-fields"
TEMPLATE_ROWS = 500
MAX_IMPORT_ROWS = 5000
MAX_UPLOAD_SIZE_BYTES = 10 * 1024 * 1024
START_ROW = 2
UNKNOWN_LABEL = "DESCONOCIDO"
OIL_SHEET = "Aceites"
GREASE_SHEET = "Grasas"

COMMON_HEADERS = [
    ("numero", "Nro"),
    ("fecha_toma", "Fecha toma"),
    ("condicion", "Condición"),
    ("fabricante", "Fabricante"),
    ("referencia_marca", "Referencia / marca"),
    ("referencia_equipo", "Máquina"),
    ("equipo_placa", "Placa manual"),
    ("periodo_servicio_aceite", "Periodo aceite"),
    ("unidad_periodo_aceite", "Unidad aceite"),
    ("periodo_servicio_equipo", "Periodo equipo"),
    ("unidad_periodo_equipo", "Unidad equipo"),
    ("observaciones", "Observaciones"),
]

def technical_fields_by_type():
    fields = list(
        CampoTecnicoMuestra.objects.filter(
            activo=True,
            deleted_at__isnull=True,
            visible_en_ingreso=True,
            tipo_muestra__in=["aceite", "grasa", "ambos"],
            catalogo__activo=True,
            catalogo__deleted_at__isnull=True,
        )
        .select_related("catalogo")
        .prefetch_related("catalogo__items")
        .order_by("orden", "nombre_visible", "id")
    )
    result = {"aceite": [], "grasa": []}
    for field in fields:
        if field.tipo_muestra in ["aceite", "ambos"]:
            result["aceite"].append(field)
        if field.tipo_muestra in ["grasa", "ambos"]:
            result["grasa"].append(field)
    return result


def technical_column_key(field):
    return f"tecnico_{field.pk}"


def template_columns(sample_type, fields_by_type=None):
    fields_by_type = fields_by_type or technical_fields_by_type()
    return COMMON_HEADERS + [
        (technical_column_key(field), field.nombre_visible)
        for field in fields_by_type[sample_type]
    ]


def option_label(obj):
    return f"{obj.pk} | {obj.nombre}"


def machine_label(obj):
    label = getattr(obj, "nombre", None) or getattr(obj, "codigo", None) or str(obj.pk)
    return f"{obj.pk} | {label}"


def parse_id(value):
    text = str(value or "").strip()
    if not text:
        return None
    if text.upper() == UNKNOWN_LABEL:
        return UNKNOWN_LABEL
    match = re.match(r"^([^|]+)\s*\|", text)
    return match.group(1).strip() if match else text


def normalize_choice(value):
    return str(value or "").strip().lower()


def as_float(value):
    if value in [None, ""]:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def add_list_validation(ws, cells, formula):
    validation = DataValidation(type="list", formula1=f"={formula}", allow_blank=True)
    validation.error = "Seleccione una opcion de la lista."
    validation.errorTitle = "Valor invalido"
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(cells)


def write_catalog_sheet(workbook, fields_by_type):
    ws = workbook.create_sheet("_CATALOGOS")
    ws.sheet_state = "veryHidden"
    columns = [
        ("Condiciones", ["usada", "nueva"]),
        ("Unidades", ["horas", "km", "millas", "dias"]),
        ("Maquinas", [machine_label(item) for item in Maquina.objects.all().order_by("id")[:3000]]),
    ]

    seen_fields = set()
    for fields in fields_by_type.values():
        for field in fields:
            if field.pk in seen_fields:
                continue
            seen_fields.add(field.pk)
            catalog = field.catalogo
            values = [UNKNOWN_LABEL]
            values += [
                option_label(item)
                for item in catalog.items.all()
                if item.activo and not item.deleted_at
            ]
            columns.append((technical_column_key(field), values))

    ranges = {}
    for column, (title, values) in enumerate(columns, start=1):
        ws.cell(row=1, column=column, value=title)
        for row, value in enumerate(values, start=2):
            ws.cell(row=row, column=column, value=value)
        letter = ws.cell(row=1, column=column).column_letter
        last_row = max(2, len(values) + 1)
        ranges[title] = f"'_CATALOGOS'!${letter}$2:${letter}${last_row}"
    ws.protection.sheet = True
    return ranges


def style_sample_sheet(ws, columns, ranges, sample_type, fields_by_type):
    fill = PatternFill("solid", fgColor="B91C1C")
    font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="404040")

    for column, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.fill, cell.font = fill, font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = max(16, min(32, len(str(header)) + 4))

    for row in range(START_ROW, TEMPLATE_ROWS + 2):
        ws.cell(row=row, column=1, value=row - 1)
        for column in range(1, len(columns) + 1):
            cell = ws.cell(row=row, column=column)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.protection = Protection(locked=column == 1)

    header_index = {key: index for index, (key, _) in enumerate(columns, start=1)}
    add_list_validation(ws, f"{ws.cell(1, header_index['condicion']).column_letter}2:{ws.cell(1, header_index['condicion']).column_letter}{TEMPLATE_ROWS + 1}", ranges["Condiciones"])
    add_list_validation(ws, f"{ws.cell(1, header_index['referencia_equipo']).column_letter}2:{ws.cell(1, header_index['referencia_equipo']).column_letter}{TEMPLATE_ROWS + 1}", ranges["Maquinas"])
    for key in ["unidad_periodo_aceite", "unidad_periodo_equipo"]:
        letter = ws.cell(1, header_index[key]).column_letter
        add_list_validation(ws, f"{letter}2:{letter}{TEMPLATE_ROWS + 1}", ranges["Unidades"])

    for field in fields_by_type[sample_type]:
        field_key = technical_column_key(field)
        letter = ws.cell(1, header_index[field_key]).column_letter
        range_key = field_key
        add_list_validation(ws, f"{letter}2:{letter}{TEMPLATE_ROWS + 1}", ranges[range_key])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(columns)).column_letter}{TEMPLATE_ROWS + 1}"
    ws.protection.sheet = True


def generate_template_workbook(token):
    fields_by_type = technical_fields_by_type()
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    readme = workbook.create_sheet("Instrucciones")
    readme["A1"] = "Plantilla de ingreso masivo de muestras"
    readme["A1"].font = Font(bold=True, size=15)
    readme["A3"] = "Use la hoja Aceites para muestras de aceite y la hoja Grasas para muestras de grasa."
    readme["A4"] = "Las columnas técnicas reflejan la configuración vigente de Campos técnicos de muestra."
    readme["A5"] = "Si el cliente no conoce un dato y el campo lo permite, seleccione DESCONOCIDO."
    readme["A6"] = "Para una muestra usada indique una máquina registrada o una placa manual, nunca ambas."
    readme.column_dimensions["A"].width = 130

    ranges = write_catalog_sheet(workbook, fields_by_type)

    sheets_meta = []
    for sample_type, sheet_name in [("aceite", OIL_SHEET), ("grasa", GREASE_SHEET)]:
        ws = workbook.create_sheet(sheet_name)
        columns = template_columns(sample_type, fields_by_type)
        style_sample_sheet(ws, columns, ranges, sample_type, fields_by_type)
        sheets_meta.append(f"{sheet_name}:{'|'.join(header for _, header in columns)}")

    meta = workbook.create_sheet("_META")
    meta.sheet_state = "veryHidden"
    meta["A1"], meta["B1"] = "schema_version", SCHEMA_VERSION
    meta["A2"], meta["B2"] = "token", str(token.token)
    meta["A3"], meta["B3"] = "template_rows", TEMPLATE_ROWS
    meta["A4"], meta["B4"] = "headers", "||".join(sheets_meta)
    meta.protection.sheet = True

    return workbook


def get_meta(workbook):
    if "_META" not in workbook.sheetnames:
        raise ValueError("La plantilla no contiene hoja de validacion.")
    meta = workbook["_META"]
    return meta["B1"].value, meta["B2"].value, int(meta["B3"].value or TEMPLATE_ROWS), str(meta["B4"].value or "")


def expected_header_meta(fields_by_type=None):
    fields_by_type = fields_by_type or technical_fields_by_type()
    return "||".join(
        f"{sheet_name}:{'|'.join(header for _, header in template_columns(sample_type, fields_by_type))}"
        for sample_type, sheet_name in [("aceite", OIL_SHEET), ("grasa", GREASE_SHEET)]
    )


def row_is_empty(ws, row, width):
    return all(ws.cell(row=row, column=column).value in [None, ""] for column in range(2, width + 1))


def validate_uploaded_workbook_file(uploaded):
    filename = str(getattr(uploaded, "name", "") or "").lower()
    if not filename.endswith(".xlsx"):
        return "Solo se aceptan archivos .xlsx generados por el sistema."
    if getattr(uploaded, "size", 0) and uploaded.size > MAX_UPLOAD_SIZE_BYTES:
        size_mb = MAX_UPLOAD_SIZE_BYTES // (1024 * 1024)
        return f"El archivo supera el tamaño máximo permitido de {size_mb} MB."
    return None


def parse_row(ws, row, sample_type, fields_by_type):
    columns = template_columns(sample_type, fields_by_type)
    raw = {key: ws.cell(row=row, column=index).value for index, (key, _) in enumerate(columns, start=1)}
    errors = []
    condition = normalize_choice(raw["condicion"])
    if condition not in ["usada", "nueva"]:
        errors.append("Condición debe ser usada o nueva.")

    data = {
        "fecha_toma": raw["fecha_toma"],
        "tipo_muestra": sample_type,
        "condicion": condition,
        "fabricante": str(raw["fabricante"] or "").strip(),
        "referencia_marca": str(raw["referencia_marca"] or "").strip(),
        "equipo_placa": str(raw["equipo_placa"] or "").strip(),
        "periodo_servicio_aceite": as_float(raw["periodo_servicio_aceite"]),
        "unidad_periodo_aceite": normalize_choice(raw["unidad_periodo_aceite"]) or "horas",
        "periodo_servicio_equipo": as_float(raw["periodo_servicio_equipo"]),
        "unidad_periodo_equipo": normalize_choice(raw["unidad_periodo_equipo"]) or "horas",
        "observaciones": str(raw["observaciones"] or "").strip(),
        "referencia_equipo": None,
        "campos_adicionales": {"fabricante": str(raw["fabricante"] or "").strip() or None},
        "atributos_tecnicos": {},
    }

    if not raw["fecha_toma"]:
        errors.append("Fecha toma es obligatoria.")
    if not data["referencia_marca"]:
        errors.append("Referencia / marca es obligatoria.")

    machine = parse_id(raw["referencia_equipo"])
    if machine and data["equipo_placa"]:
        errors.append("Seleccione una máquina registrada o una placa manual, no ambas.")
    if machine:
        if Maquina.objects.filter(pk=machine).exists():
            data["referencia_equipo"] = str(machine)
        else:
            errors.append(f"Máquina inválida ({machine}).")
    if condition == "usada" and not data["referencia_equipo"] and not data["equipo_placa"]:
        errors.append("Muestra usada requiere Máquina o Placa manual.")

    for field in fields_by_type[sample_type]:
        field_key = technical_column_key(field)
        label = field.nombre_visible
        catalog = field.catalogo
        raw_value = raw.get(field_key)
        item_id = parse_id(raw_value)
        unknown = item_id == UNKNOWN_LABEL

        active_items = catalog.items.filter(activo=True, deleted_at__isnull=True)
        allows_unknown = field.permite_desconocido and catalog.permite_desconocido
        if not active_items.exists() and allows_unknown and not item_id:
            unknown = True

        item = None if unknown else active_items.filter(pk=item_id).first()
        if item_id and not unknown and not item:
            errors.append(f"{label}: opción inválida ({item_id}).")
        if unknown and not allows_unknown:
            errors.append(f"{label}: el campo no permite DESCONOCIDO.")
        if field.obligatorio and not item and not unknown:
            errors.append(f"{label}: seleccione una opción.")

        data["atributos_tecnicos"][str(catalog.id)] = {
            "item": str(item.id) if item else "",
            "desconocido": unknown,
            "campo_tecnico_muestra": str(field.id),
        }
        data[field.codigo] = str(item.id) if item else ""
        data[f"{field.codigo}_desconocido"] = unknown

    return data, errors


class SampleBatchExcelImportViewSet(ActionPermissionMixin, viewsets.ViewSet):
    permission_classes = [IsAuthenticated, DenyReadOnlyWrite, HasSecurityPermission]
    permission_action_map = {
        "template": "muestras.descargar_plantilla_excel",
        "preview": "muestras.previsualizar_excel",
    }
    parser_classes = [MultiPartParser, FormParser]

    @action(detail=False, methods=["get"], url_path="template")
    def template(self, request):
        cliente_empresa = request.query_params.get("cliente_empresa") or None
        if cliente_empresa and not (request.user.is_superuser or request.user.role == User.Role.GLOBAL):
            if str(cliente_empresa) != str(request.user.empresa_id):
                return Response({"detail": "No puede generar plantilla para otra empresa."}, status=status.HTTP_403_FORBIDDEN)
        token = SampleBatchExcelTemplateToken.objects.create(
            requested_by=request.user,
            cliente_empresa_id=cliente_empresa,
            schema_version=SCHEMA_VERSION,
            expires_at=timezone.now() + timedelta(hours=8),
        )
        filename = f"MUESTRAS_LOTE_{SCHEMA_VERSION}_{str(token.token).split('-')[0]}.xlsx"
        token.filename = filename
        token.save(update_fields=["filename"])
        workbook = generate_template_workbook(token)
        buffer = BytesIO()
        workbook.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @action(detail=False, methods=["post"], url_path="preview")
    def preview(self, request):
        uploaded = request.FILES.get("file")
        if not uploaded:
            return Response({"detail": "Debe subir un archivo Excel."}, status=status.HTTP_400_BAD_REQUEST)
        file_error = validate_uploaded_workbook_file(uploaded)
        if file_error:
            return Response({"detail": file_error}, status=status.HTTP_400_BAD_REQUEST)
        try:
            workbook = load_workbook(uploaded, data_only=True)
            schema, token_value, template_rows, stored_headers = get_meta(workbook)
            token = SampleBatchExcelTemplateToken.objects.get(token=uuid.UUID(str(token_value)), schema_version=SCHEMA_VERSION)
        except Exception as exc:
            return Response({"detail": f"Archivo invalido: {exc}"}, status=status.HTTP_400_BAD_REQUEST)
        if schema != SCHEMA_VERSION or not token.is_valid:
            return Response({"detail": "Plantilla vencida o incompatible. Genere una nueva."}, status=status.HTTP_400_BAD_REQUEST)
        if token.requested_by_id and token.requested_by_id != request.user.id:
            return Response({"detail": "La plantilla pertenece a otro usuario. Genere una nueva."}, status=status.HTTP_403_FORBIDDEN)

        fields_by_type = technical_fields_by_type()
        if stored_headers != expected_header_meta(fields_by_type):
            return Response({"detail": "La configuración técnica cambió. Genere una nueva plantilla."}, status=status.HTTP_400_BAD_REQUEST)

        rows = []
        truncated = False
        for sample_type, sheet_name in [("aceite", OIL_SHEET), ("grasa", GREASE_SHEET)]:
            if sheet_name not in workbook.sheetnames:
                return Response({"detail": f"La plantilla no contiene la hoja {sheet_name}."}, status=status.HTTP_400_BAD_REQUEST)
            ws = workbook[sheet_name]
            columns = template_columns(sample_type, fields_by_type)
            headers = [header for _, header in columns]
            if [ws.cell(1, column).value for column in range(1, len(headers) + 1)] != headers:
                return Response({"detail": f"La estructura de columnas de {sheet_name} fue modificada."}, status=status.HTTP_400_BAD_REQUEST)

            last_candidate_row = max(ws.max_row, int(template_rows or TEMPLATE_ROWS) + 1)
            last_row = min(last_candidate_row, MAX_IMPORT_ROWS + 1)
            if last_candidate_row > MAX_IMPORT_ROWS + 1:
                truncated = True

            for row in range(START_ROW, last_row + 1):
                if len(rows) >= MAX_IMPORT_ROWS:
                    truncated = True
                    break
                if row_is_empty(ws, row, len(headers)):
                    continue
                data, errors = parse_row(ws, row, sample_type, fields_by_type)
                rows.append({"sheet": sheet_name, "row_number": row, "is_valid": not errors, "errors": errors, "data": data})
            if len(rows) >= MAX_IMPORT_ROWS:
                break

        return Response({
            "template_token": str(token.token),
            "schema_version": SCHEMA_VERSION,
            "total_rows": len(rows),
            "valid_rows": sum(1 for row in rows if row["is_valid"]),
            "error_rows": sum(1 for row in rows if not row["is_valid"]),
            "total_errors": sum(len(row["errors"]) for row in rows),
            "template_rows": template_rows,
            "max_import_rows": MAX_IMPORT_ROWS,
            "truncated": truncated,
            "rows": rows,
        })
