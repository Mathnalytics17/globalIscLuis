import json
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from django.db.models import Count, Max
from django.utils.text import slugify

from django.db import transaction
from django.utils.text import slugify
from django.utils import timezone
from rest_framework import filters, status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from django.http import HttpResponse
from apps.utils.pagination import StandardResultsSetPagination
from permissions import ActionPermissionMixin, DenyReadOnlyWrite, HasSecurityPermission

from apps.misc.api.models.dynamicTechnicalConfig.index import (
    CatalogoTecnico,
    CatalogoTecnicoVersion,
    CampoTecnicoMuestra,
    CatalogoTecnicoCampo,
    CatalogoTecnicoItem,
    CatalogoTecnicoItemValor,
    EscalaComparacion,
    EscalaComparacionItem,
    PruebaFuenteLimite,
    PruebaFuenteCatalogo,
    CriterioEvaluacionLimite,
    CriterioCatalogoSeleccion,
    PruebaLimiteCampo,
)
from apps.misc.api.models.pruebas.index import Prueba, PruebaResultado, PruebaResultadoDivision, PruebaResultadoComponente
from apps.misc.api.serializers.dynamicTechnicalConfig.index import (
    CatalogoTecnicoCampoSerializer,
    CatalogoTecnicoItemSerializer,
    CatalogoTecnicoItemValorSerializer,
    CatalogoTecnicoSerializer,
    CatalogoTecnicoVersionSerializer,
    CampoTecnicoMuestraSerializer,
    EscalaComparacionItemSerializer,
    EscalaComparacionSerializer,
    PruebaFuenteLimiteSerializer,
    CriterioEvaluacionLimiteSerializer,
    PruebaLimiteCampoSerializer,
)
from apps.misc.api.services.limit_contract import (
    comparison_from_component as _comparison_from_component,
    test_limit_leaves as _test_limit_leaves,
)
from apps.muestras.api.models.muestras.index import Muestra
from apps.muestras.api.services.limit_engine import resolve_limit_for_result, resolve_and_evaluate_result


class SoftDeleteViewSet(ActionPermissionMixin, viewsets.ModelViewSet):
    pagination_class = StandardResultsSetPagination
    permission_classes = [IsAuthenticated, DenyReadOnlyWrite, HasSecurityPermission]
    default_read_permission = "config_tecnica.ver"
    default_write_permission = "config_tecnica.editar"
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.action != "restore" and self.request.query_params.get("incluir_eliminados") != "true":
            queryset = queryset.filter(deleted_at__isnull=True)
        activo = self.request.query_params.get("activo")
        if activo is not None and hasattr(queryset.model, "activo"):
            normalized = str(activo).strip().lower()
            if normalized in {"true", "1", "yes", "si"}:
                queryset = queryset.filter(activo=True)
            elif normalized in {"false", "0", "no"}:
                queryset = queryset.filter(activo=False)
        return queryset

    def perform_destroy(self, instance):
        instance.soft_delete()

    @action(detail=True, methods=["post"])
    def restore(self, request, pk=None):
        base_queryset = self.queryset if self.queryset is not None else self.get_queryset()
        instance = base_queryset.get(pk=pk)
        instance.restore()
        return Response(self.get_serializer(instance).data, status=status.HTTP_200_OK)


class CatalogoTecnicoViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "catalogos_tecnicos.ver",
        "retrieve": "catalogos_tecnicos.ver",
        "sample_form": "catalogos_tecnicos.ver",
        "create": "catalogos_tecnicos.crear",
        "update": "catalogos_tecnicos.editar",
        "partial_update": "catalogos_tecnicos.editar",
        "destroy": "catalogos_tecnicos.eliminar",
        "restore": "catalogos_tecnicos.restaurar",
    }
    serializer_class = CatalogoTecnicoSerializer
    search_fields = ["nombre", "codigo", "descripcion"]
    ordering_fields = ["orden", "nombre", "codigo", "created_at"]

    def get_queryset(self):
        self.queryset = CatalogoTecnico.objects.prefetch_related("campos", "versiones", "items__valores")
        queryset = super().get_queryset()
        tipo_muestra = self.request.query_params.get("tipo_muestra")
        requerido = self.request.query_params.get("es_requerido_en_muestra")
        if tipo_muestra:
            queryset = queryset.filter(tipo_muestra__in=[tipo_muestra, "ambos"])
        if requerido is not None:
            queryset = queryset.filter(es_requerido_en_muestra=requerido.lower() == "true")
        return queryset

    @transaction.atomic
    def perform_create(self, serializer):
        catalogo = serializer.save()
        version = CatalogoTecnicoVersion.objects.create(catalogo=catalogo, numero=1, nombre="Versión inicial")
        catalogo.version_actual = version
        catalogo.save(update_fields=["version_actual", "updated_at"])

    @action(detail=False, methods=["get"], url_path="sample-form")
    def sample_form(self, request):
        tipo = request.query_params.get("tipo_muestra")
        fields = CampoTecnicoMuestra.objects.filter(
            activo=True,
            deleted_at__isnull=True,
            visible_en_ingreso=True,
        ).select_related("catalogo").prefetch_related("catalogo__campos", "catalogo__items__valores")
        if tipo:
            fields = fields.filter(
                tipo_muestra__in=[tipo, "ambos"],
                catalogo__tipo_muestra__in=[tipo, "ambos"],
            )
        fields = fields.order_by("tipo_muestra", "orden", "nombre_visible")
        return Response(CampoTecnicoMuestraSerializer(fields, many=True).data)


class CatalogoTecnicoVersionViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "catalogos_tecnicos.ver", "retrieve": "catalogos_tecnicos.ver",
        "create": "catalogos_tecnicos.crear", "update": "catalogos_tecnicos.editar",
        "partial_update": "catalogos_tecnicos.editar", "destroy": "catalogos_tecnicos.eliminar",
        "restore": "catalogos_tecnicos.restaurar",
    }
    serializer_class = CatalogoTecnicoVersionSerializer
    search_fields = ["nombre", "norma_referencia", "notas"]
    ordering_fields = ["numero", "fecha_vigencia", "created_at"]

    def get_queryset(self):
        self.queryset = CatalogoTecnicoVersion.objects.annotate(items_count=Count("items"))
        queryset = super().get_queryset()
        catalogo = self.request.query_params.get("catalogo")
        return queryset.filter(catalogo_id=catalogo) if catalogo else queryset

    @transaction.atomic
    def perform_create(self, serializer):
        catalogo = serializer.validated_data["catalogo"]
        numero = serializer.validated_data.get("numero")
        if not numero:
            numero = (CatalogoTecnicoVersion.objects.filter(catalogo=catalogo).aggregate(maximo=Max("numero"))["maximo"] or 0) + 1
        version = serializer.save(numero=numero)
        catalogo.version_actual = version
        catalogo.save(update_fields=["version_actual", "updated_at"])

    @transaction.atomic
    def perform_update(self, serializer):
        version = serializer.save()
        if self.request.data.get("hacer_actual") in [True, "true", "True", "1"]:
            version.catalogo.version_actual = version
            version.catalogo.save(update_fields=["version_actual", "updated_at"])


class CampoTecnicoMuestraViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "catalogos_tecnicos.ver",
        "retrieve": "catalogos_tecnicos.ver",
        "create": "catalogos_tecnicos.editar",
        "update": "catalogos_tecnicos.editar",
        "partial_update": "catalogos_tecnicos.editar",
        "destroy": "catalogos_tecnicos.eliminar",
        "restore": "catalogos_tecnicos.restaurar",
    }
    queryset = CampoTecnicoMuestra.objects.select_related("catalogo").prefetch_related("catalogo__items", "catalogo__campos")
    serializer_class = CampoTecnicoMuestraSerializer
    search_fields = ["nombre_visible", "codigo", "catalogo__nombre"]
    ordering_fields = ["tipo_muestra", "orden", "nombre_visible", "created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        tipo_muestra = self.request.query_params.get("tipo_muestra")
        visible_en_ingreso = self.request.query_params.get("visible_en_ingreso")
        visible_en_asignacion = self.request.query_params.get("visible_en_asignacion")
        if tipo_muestra:
            queryset = queryset.filter(tipo_muestra__in=[tipo_muestra, "ambos"])
        if visible_en_ingreso is not None:
            queryset = queryset.filter(visible_en_ingreso=visible_en_ingreso.lower() == "true")
        if visible_en_asignacion is not None:
            queryset = queryset.filter(visible_en_asignacion=visible_en_asignacion.lower() == "true")
        return queryset


class CatalogoTecnicoCampoViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "catalogos_tecnicos.ver",
        "retrieve": "catalogos_tecnicos.ver",
        "create": "catalogos_tecnicos.editar",
        "update": "catalogos_tecnicos.editar",
        "partial_update": "catalogos_tecnicos.editar",
        "destroy": "catalogos_tecnicos.eliminar",
        "restore": "catalogos_tecnicos.restaurar",
    }
    queryset = CatalogoTecnicoCampo.objects.select_related("catalogo")
    serializer_class = CatalogoTecnicoCampoSerializer
    search_fields = ["nombre", "codigo", "catalogo__nombre"]
    ordering_fields = ["orden", "nombre", "codigo"]

    def get_queryset(self):
        queryset = super().get_queryset()
        catalogo = self.request.query_params.get("catalogo")
        return queryset.filter(catalogo_id=catalogo) if catalogo else queryset


class CatalogoTecnicoItemViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "catalogos_tecnicos.ver",
        "retrieve": "catalogos_tecnicos.ver",
        "create": "catalogos_tecnicos.crear",
        "update": "catalogos_tecnicos.editar",
        "partial_update": "catalogos_tecnicos.editar",
        "destroy": "catalogos_tecnicos.eliminar",
        "restore": "catalogos_tecnicos.restaurar",
        "import_excel": "catalogos_tecnicos.crear",
        "template_excel": "catalogos_tecnicos.ver",
    }
    queryset = CatalogoTecnicoItem.objects.select_related("catalogo", "version").prefetch_related("valores__campo")
    serializer_class = CatalogoTecnicoItemSerializer
    search_fields = ["nombre", "codigo", "catalogo__nombre"]
    ordering_fields = ["nombre", "codigo", "created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        catalogo = self.request.query_params.get("catalogo")
        version = self.request.query_params.get("version")
        catalogo_codigo = self.request.query_params.get("catalogo_codigo")
        if catalogo:
            queryset = queryset.filter(catalogo_id=catalogo)
        if version:
            queryset = queryset.filter(version_id=version)
        if catalogo_codigo:
            queryset = queryset.filter(catalogo__codigo=catalogo_codigo)
        return queryset.order_by("nombre", "id")

    @action(detail=False, methods=["post"], url_path="import-excel")
    @transaction.atomic
    def import_excel(self, request):
        version_id = request.data.get("version")
        upload = request.FILES.get("archivo")
        if not version_id or not upload:
            return Response({"detail": "Seleccione la versión y un archivo Excel."}, status=status.HTTP_400_BAD_REQUEST)
        if not upload.name.lower().endswith(".xlsx"):
            return Response({"detail": "El archivo debe tener extensión .xlsx."}, status=status.HTTP_400_BAD_REQUEST)
        try:
            version = CatalogoTecnicoVersion.objects.select_related("catalogo").get(pk=version_id, deleted_at__isnull=True)
            workbook = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True)
            sheet = workbook.active
        except (CatalogoTecnicoVersion.DoesNotExist, ValueError, OSError) as exc:
            return Response({"detail": f"No se pudo leer el archivo o la versión: {exc}"}, status=status.HTTP_400_BAD_REQUEST)

        headers = [str(cell.value or "").strip().lower() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        required = {"nombre", "codigo", "descripcion"}
        if not required.issubset(set(headers)):
            return Response({"detail": "La primera fila debe contener: nombre, codigo, descripcion."}, status=status.HTTP_400_BAD_REQUEST)
        positions = {name: headers.index(name) for name in required}
        created = updated = 0
        errors = []
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            nombre = str(row[positions["nombre"]] or "").strip()
            codigo = slugify(str(row[positions["codigo"]] or nombre)).replace("-", "_")
            descripcion = str(row[positions["descripcion"]] or "").strip()
            if not nombre:
                if any(value not in (None, "") for value in row): errors.append({"fila": number, "error": "El nombre es obligatorio."})
                continue
            if not codigo:
                errors.append({"fila": number, "error": "El código no es válido."})
                continue
            item, is_created = CatalogoTecnicoItem.objects.update_or_create(
                version=version, codigo=codigo,
                defaults={"catalogo": version.catalogo, "nombre": nombre, "descripcion": descripcion, "activo": True, "deleted_at": None},
            )
            created += int(is_created)
            updated += int(not is_created)
        if errors:
            transaction.set_rollback(True)
            return Response({"detail": "La importación no se aplicó.", "errores": errors[:50]}, status=status.HTTP_400_BAD_REQUEST)
        return Response({"creados": created, "actualizados": updated, "total": created + updated})

    @action(detail=False, methods=["get"], url_path="template-excel")
    def template_excel(self, request):
        """Genera la plantilla oficial para evitar archivos con columnas ambiguas."""
        version_id = request.query_params.get("version")
        if not version_id:
            return Response({"detail": "Seleccione una versión para descargar la plantilla."}, status=status.HTTP_400_BAD_REQUEST)
        try:
            version = CatalogoTecnicoVersion.objects.select_related("catalogo").get(pk=version_id, deleted_at__isnull=True)
        except CatalogoTecnicoVersion.DoesNotExist:
            return Response({"detail": "La versión seleccionada no existe."}, status=status.HTTP_404_NOT_FOUND)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ítems"
        headers = ["nombre", "codigo", "descripcion"]
        sheet.append(headers)
        sheet.append(["Ejemplo de ítem", "ejemplo_item", "Descripción opcional del ítem"])
        header_fill = PatternFill("solid", fgColor="E9232D")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 28
        sheet.column_dimensions["C"].width = 56
        stream = BytesIO()
        workbook.save(stream)
        filename = f"plantilla_{slugify(version.catalogo.codigo)}_v{version.numero}.xlsx"
        response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class CatalogoTecnicoItemValorViewSet(ActionPermissionMixin, viewsets.ModelViewSet):
    pagination_class = StandardResultsSetPagination
    permission_classes = [IsAuthenticated, DenyReadOnlyWrite, HasSecurityPermission]
    default_read_permission = "catalogos_tecnicos.ver"
    default_write_permission = "catalogos_tecnicos.editar"
    serializer_class = CatalogoTecnicoItemValorSerializer

    def get_queryset(self):
        queryset = CatalogoTecnicoItemValor.objects.select_related("item", "campo", "item__catalogo")
        catalogo = self.request.query_params.get("catalogo")
        item = self.request.query_params.get("item")
        campo = self.request.query_params.get("campo")
        if catalogo:
            queryset = queryset.filter(item__catalogo_id=catalogo)
        if item:
            queryset = queryset.filter(item_id=item)
        if campo:
            queryset = queryset.filter(campo_id=campo)
        return queryset.order_by("item__nombre", "campo__orden", "id")


class EscalaComparacionViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "config_tecnica.ver",
        "retrieve": "config_tecnica.ver",
        "create": "config_tecnica.editar",
        "update": "config_tecnica.editar",
        "partial_update": "config_tecnica.editar",
        "destroy": "config_tecnica.editar",
        "restore": "config_tecnica.editar",
    }
    queryset = EscalaComparacion.objects.prefetch_related("items")
    serializer_class = EscalaComparacionSerializer
    search_fields = ["nombre", "codigo", "descripcion"]
    ordering_fields = ["nombre", "codigo", "created_at"]


class EscalaComparacionItemViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "config_tecnica.ver",
        "retrieve": "config_tecnica.ver",
        "create": "config_tecnica.editar",
        "update": "config_tecnica.editar",
        "partial_update": "config_tecnica.editar",
        "reorder": "config_tecnica.editar",
        "destroy": "config_tecnica.editar",
        "restore": "config_tecnica.editar",
    }
    queryset = EscalaComparacionItem.objects.select_related("escala")
    serializer_class = EscalaComparacionItemSerializer
    search_fields = ["etiqueta", "valor_normalizado", "escala__nombre"]
    ordering_fields = ["orden", "etiqueta", "created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        escala = self.request.query_params.get("escala")
        return queryset.filter(escala_id=escala) if escala else queryset

    @action(detail=False, methods=["post"], url_path="reorder")
    @transaction.atomic
    def reorder(self, request):
        escala_id = request.data.get("escala")
        items = request.data.get("items") or []
        if not escala_id:
            return Response({"escala": ["Este campo es requerido."]}, status=status.HTTP_400_BAD_REQUEST)
        if not isinstance(items, list) or not items:
            return Response({"items": ["Debe enviar una lista de items."]}, status=status.HTTP_400_BAD_REQUEST)

        existing = {
            item.id: item
            for item in EscalaComparacionItem.objects.filter(
                escala_id=escala_id,
                deleted_at__isnull=True,
            )
        }
        requested_ids = []
        for raw in items:
            try:
                item_id = int(raw.get("id"))
            except (TypeError, ValueError, AttributeError):
                return Response({"items": ["Cada item debe incluir un id valido."]}, status=status.HTTP_400_BAD_REQUEST)
            if item_id not in existing:
                return Response({"items": [f"El item {item_id} no pertenece a esta escala."]}, status=status.HTTP_400_BAD_REQUEST)
            requested_ids.append(item_id)

        ordered_ids = requested_ids + [item_id for item_id in existing.keys() if item_id not in requested_ids]
        for index, item_id in enumerate(ordered_ids, start=1):
            item = existing[item_id]
            if item.orden != index:
                item.orden = index
                item.save(update_fields=["orden", "updated_at"])

        queryset = EscalaComparacionItem.objects.filter(escala_id=escala_id, deleted_at__isnull=True).order_by("orden", "id")
        return Response(self.get_serializer(queryset, many=True).data, status=status.HTTP_200_OK)


class PruebaFuenteLimiteViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "valores_limite.ver",
        "retrieve": "valores_limite.ver",
        "resolve_preview": "valores_limite.ver",
        "create": "valores_limite.editar_matriz",
        "update": "valores_limite.editar_matriz",
        "partial_update": "valores_limite.editar_matriz",
        "destroy": "valores_limite.editar_matriz",
        "restore": "valores_limite.editar_matriz",
        "generate_fields": "valores_limite.editar_matriz",
        "configure": "valores_limite.editar_matriz",
        "matrix_template": "valores_limite.ver",
        "import_matrix": "valores_limite.editar_matriz",
    }
    queryset = PruebaFuenteLimite.objects.select_related("prueba", "catalogo_fuente", "campo_tecnico_muestra", "campo_tecnico_muestra__catalogo").prefetch_related(
        "catalogos_configurados__catalogo",
        "catalogos_configurados__campo_tecnico_muestra",
        "criterios__selecciones_catalogo__catalogo",
        "criterios__selecciones_catalogo__item",
        "campos_limite__resultado",
        "campos_limite__division",
        "campos_limite__componente",
    )
    serializer_class = PruebaFuenteLimiteSerializer
    search_fields = ["prueba__nombre_variable", "prueba__acronimo", "catalogo_fuente__nombre"]
    ordering_fields = ["prioridad", "prueba", "created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        prueba = self.request.query_params.get("prueba")
        catalogo = self.request.query_params.get("catalogo_fuente")
        if prueba:
            queryset = queryset.filter(prueba_id=prueba)
        if catalogo:
            queryset = queryset.filter(catalogo_fuente_id=catalogo)
        return queryset

    @action(detail=True, methods=["post"], url_path="generate-fields")
    @transaction.atomic
    def generate_fields(self, request, pk=None):
        fuente = self.get_object()
        operador = request.data.get("operador") or "max"
        unidad = request.data.get("unidad") or ""
        fields = _build_limit_fields_from_prueba(
            fuente,
            operador=operador,
            unidad=unidad,
            tipo_comparacion=request.data.get("tipo_comparacion") or "numerica",
            escala_comparacion_id=request.data.get("escala_comparacion") or None,
            modo_ordinal=request.data.get("modo_ordinal") or "orden",
            evaluacion_opciones=request.data.get("evaluacion_opciones"),
        )
        return Response(PruebaLimiteCampoSerializer(fields, many=True).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["get"], url_path="matrix-template")
    def matrix_template(self, request, pk=None):
        source = self.get_object()
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Matriz de limites"
        headers = ["criterio_codigo", "campo_codigo", "operacion", "usa_semaforo", "min_critico", "min_aceptable", "max_aceptable", "max_critico", "valor_objetivo", "valores_permitidos", "comentario"]
        sheet.append(headers)
        for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="E9232D")
        fields = (source.configuracion_regla or {}).get("campos", [])
        for criterion in source.criterios.filter(activo=True, deleted_at__isnull=True):
            for field in fields:
                if field.get("origen_limite") != "catalogo": continue
                rule = (criterion.valores_limite or {}).get(field.get("codigo"), {})
                if not isinstance(rule, dict): rule = {"valor_esperado": rule}
                sheet.append([criterion.codigo, field.get("codigo"), rule.get("operador", field.get("operador", "max")), "SI" if rule.get("usar_amarillo") else "NO", rule.get("min_critico", "NA"), rule.get("min_aceptable", "NA"), rule.get("max_aceptable", "NA"), rule.get("max_critico", "NA"), rule.get("valor_esperado", "NA"), "NA", ""])
        sheet.freeze_panes = "A2"
        for col, width in zip("ABCDEFGHIJK", [26, 28, 16, 16, 18, 18, 18, 18, 22, 28, 38]): sheet.column_dimensions[col].width = width
        stream = BytesIO(); workbook.save(stream)
        response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="matriz_limites_{source.prueba_id}.xlsx"'
        return response

    @action(detail=True, methods=["post"], url_path="import-matrix")
    @transaction.atomic
    def import_matrix(self, request, pk=None):
        source = self.get_object(); upload = request.FILES.get("archivo")
        if not upload or not upload.name.lower().endswith(".xlsx"): return Response({"detail": "Cargue la plantilla .xlsx."}, status=status.HTTP_400_BAD_REQUEST)
        try: sheet = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True).active
        except Exception: return Response({"detail": "No se pudo leer el archivo Excel."}, status=status.HTTP_400_BAD_REQUEST)
        headers = [str(c.value or "").strip() for c in next(sheet.iter_rows(min_row=1, max_row=1))]
        required = ["criterio_codigo", "campo_codigo", "operacion", "usa_semaforo", "min_critico", "min_aceptable", "max_aceptable", "max_critico", "valor_objetivo", "valores_permitidos"]
        if any(name not in headers for name in required): return Response({"detail": "La plantilla no contiene todas las columnas requeridas."}, status=status.HTTP_400_BAD_REQUEST)
        pos = {name: headers.index(name) for name in headers}; fields = {f.get("codigo"): f for f in (source.configuracion_regla or {}).get("campos", []) if f.get("origen_limite") == "catalogo"}; criteria = {c.codigo: c for c in source.criterios.filter(activo=True, deleted_at__isnull=True)}; errors=[]; updates=[]
        allowed = {"max", "min", "between", "eq", "neq", "in", "not_in", "informativo"}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            code, field_code = str(row[pos["criterio_codigo"]] or "").strip(), str(row[pos["campo_codigo"]] or "").strip()
            if not code and not field_code: continue
            operation = str(row[pos["operacion"]] or "").strip().lower()
            if code not in criteria or field_code not in fields or operation not in allowed: errors.append({"fila": row_number, "error": "Criterio, campo u operación no válida."}); continue
            val=lambda name: None if str(row[pos[name]] or "").strip().upper() in {"", "NA"} else str(row[pos[name]]).strip()
            rule={"operador": operation, "usar_amarillo": str(row[pos["usa_semaforo"]] or "").strip().upper() == "SI", "min_critico": val("min_critico"), "min_aceptable": val("min_aceptable"), "max_aceptable": val("max_aceptable"), "max_critico": val("max_critico"), "valor_esperado": val("valor_objetivo") or val("valores_permitidos")}
            needed = {"max":["max_aceptable"], "min":["min_aceptable"], "between":["min_aceptable","max_aceptable"], "eq":["valor_esperado"], "neq":["valor_esperado"], "in":["valor_esperado"], "not_in":["valor_esperado"], "informativo":[]}[operation]
            if any(rule[key] is None for key in needed): errors.append({"fila": row_number, "error": "Faltan valores requeridos para la operación."}); continue
            updates.append((criteria[code], field_code, rule))
        if errors: transaction.set_rollback(True); return Response({"detail":"La importación no se aplicó.","errores":errors[:100]}, status=status.HTTP_400_BAD_REQUEST)
        for criterion, field_code, rule in updates:
            values = criterion.valores_limite or {}; values[field_code] = rule; criterion.valores_limite = values; criterion.save(update_fields=["valores_limite", "updated_at"])
        return Response({"actualizados": len(updates)})

    @action(detail=False, methods=["post"], url_path="configure")
    @transaction.atomic
    def configure(self, request):
        """Guarda la configuración completa de evaluación de una prueba en una transacción."""
        prueba = Prueba.objects.prefetch_related(
            "resultados__divisiones__componentes"
        ).get(pk=request.data.get("prueba"))
        source_type = request.data.get("fuente") or "catalogos"
        catalogs = request.data.get("catalogos") or []
        mode = request.data.get("modo_seleccion") or "seleccionar_asignacion"
        decision = request.data.get("decision") or {}
        decision_fields = decision.get("campos") if isinstance(decision.get("campos"), list) else []
        decision["informativos"] = [
            str(item.get("codigo")) for item in decision_fields if item.get("informativo") or item.get("evalua") is False
        ]
        for legacy_key in [
            "campos_requeridos", "criticos", "pesos", "peso_minimo",
            "puntaje_minimo", "umbral_peso_fallado", "fallo_critico_rechaza",
        ]:
            decision.pop(legacy_key, None)
        decision["modo"] = "semaforo"
        decision["version"] = 2
        result_rules = request.data.get("reglas_resultados") or {}
        field_configs = request.data.get("campos") or []
        try:
            for field_config in field_configs:
                if field_config.get("origen_limite") == "catalogo":
                    continue
                _validate_semaphore_rule(
                    field_config.get("valor_global"),
                    field_config.get("nombre") or field_config.get("codigo") or "Campo",
                    field_config,
                )
            field_by_code = {
                str(item.get("codigo")): item
                for item in field_configs
                if item.get("codigo")
            }
            for criterion in request.data.get("criterios") or []:
                for field_code, rule in (criterion.get("valores") or {}).items():
                    field_config = field_by_code.get(str(field_code), {})
                    _validate_semaphore_rule(
                        rule,
                        f"{criterion.get('nombre') or 'Criterio'} / {field_config.get('nombre') or field_code}",
                        field_config,
                    )
        except ValueError as error:
            return Response({"limites": [str(error)]}, status=status.HTTP_400_BAD_REQUEST)

        has_catalog_fields = any(
            item.get("origen_limite") == "catalogo"
            and item.get("participa") is not False
            and item.get("evalua") is not False
            for item in field_configs
        )
        source_limit_type = "catalogo" if has_catalog_fields else "global"
        primary_catalog_id = catalogs[0].get("catalogo") if catalogs and has_catalog_fields else None
        source = PruebaFuenteLimite.objects.filter(prueba=prueba).order_by("id").first()
        if source is None:
            source = PruebaFuenteLimite(prueba=prueba)
        source.tipo_limite = source_limit_type
        source.catalogo_fuente_id = primary_catalog_id
        source.campo_tecnico_muestra_id = None
        source.modo_seleccion = mode
        source.politica_evaluacion = "todas_deben_cumplir"
        source.minimo_campos_cumplidos = None
        source.configuracion_regla = decision
        source.reglas_resultados = result_rules
        source.publicada = bool(request.data.get("publicar", True))
        source.activo = True
        source.deleted_at = None
        source.version = (source.version or 0) + (1 if source.pk else 0)
        source.save()

        PruebaFuenteLimite.objects.filter(prueba=prueba).exclude(pk=source.pk).update(
            activo=False, deleted_at=timezone.now()
        )
        source.catalogos_configurados.all().delete()
        for index, relation in enumerate(catalogs, start=1):
            PruebaFuenteCatalogo.objects.create(
                fuente_limite=source,
                catalogo_id=relation.get("catalogo"),
                campo_tecnico_muestra_id=relation.get("campo_tecnico_muestra") or None,
                orden=index,
            )

        source.campos_limite.all().update(activo=False, deleted_at=timezone.now())
        leaves = _test_limit_leaves(prueba)
        field_by_code = {
            str(item.get("codigo")): item
            for item in field_configs
            if item.get("codigo")
        }
        canonical_fields = []
        for leaf in leaves:
            config = field_by_code.get(str(leaf["codigo"]))
            if config is None:
                config = next((item for item in field_configs if (
                    str(item.get("componente") or "") == str(getattr(leaf.get("componente"), "pk", ""))
                    or (
                        not leaf.get("componente")
                        and str(item.get("resultado") or "") == str(getattr(leaf.get("resultado"), "pk", ""))
                    )
                )), {})
            canonical_fields.append({
                **config,
                "codigo": leaf["codigo"],
                "nombre": leaf["nombre"],
                "resultado": getattr(leaf.get("resultado"), "pk", None),
                "division": getattr(leaf.get("division"), "pk", None),
                "componente": getattr(leaf.get("componente"), "pk", None),
                "unidad": (
                    getattr(leaf.get("resultado"), "unidad_medida", "")
                    or config.get("unidad")
                    or ""
                ),
            })
        field_configs = canonical_fields
        decision["campos"] = canonical_fields
        source.configuracion_regla = decision
        source.save(update_fields=["configuracion_regla", "updated_at"])

        saved_fields = []
        for leaf_index, leaf in enumerate(leaves, start=1):
            config = next(
                (item for item in field_configs if str(item.get("codigo")) == str(leaf["codigo"])),
                {},
            )
            comparison = config.get("tipo_comparacion") or _comparison_from_component(leaf.get("componente"))
            origin = config.get("origen_limite") or ("informativo" if comparison == "comentario" else source_limit_type)
            if (
                comparison == "comentario"
                or origin == "informativo"
                or config.get("participa") is False
                or config.get("evalua") is False
            ):
                continue
            operator = _normalize_operator(config.get("operador") or "max")
            global_rule = config.get("valor_global") if origin in ["directo", "escala", "booleano"] else None
            saved_fields.append(_upsert_limit_field(
                fuente=source,
                nombre=leaf["nombre"],
                codigo=leaf["codigo"],
                operador=operator,
                unidad=config.get("unidad") or getattr(leaf.get("resultado"), "unidad_medida", "") or "",
                orden=len(saved_fields) + 1,
                tipo_comparacion=comparison,
                escala_comparacion_id=config.get("escala_comparacion") or getattr(leaf.get("componente"), "escala_comparacion_id", None),
                modo_ordinal="orden",
                evaluacion_opciones=config.get("evaluacion_booleano"),
                valor_global=json.dumps(global_rule, ensure_ascii=False) if isinstance(global_rule, dict) else global_rule,
                resultado=leaf.get("resultado"),
                division=leaf.get("division"),
                componente=leaf.get("componente"),
            ))

        source.criterios.all().delete()
        for criterion_index, criterion_data in enumerate(request.data.get("criterios") or [], start=1):
            criterion = CriterioEvaluacionLimite.objects.create(
                fuente_limite=source,
                nombre=criterion_data.get("nombre") or f"Criterio {criterion_index}",
                codigo=criterion_data.get("codigo") or slugify(criterion_data.get("nombre") or f"criterio-{criterion_index}").replace("-", "_"),
                tipo_criterio=criterion_data.get("tipo_criterio") or "valor_fijo",
                valor=criterion_data.get("valor") or criterion_data.get("nombre"),
                catalogo_item_id=criterion_data.get("catalogo_item") or None,
                escala_item_id=criterion_data.get("escala_item") or None,
                valores_limite=criterion_data.get("valores") or {},
                orden=criterion_index,
            )
            for selection in criterion_data.get("selecciones_catalogo") or []:
                CriterioCatalogoSeleccion.objects.create(
                    criterio=criterion,
                    catalogo_id=selection.get("catalogo"),
                    item_id=selection.get("item"),
                )

        # Una nueva versión invalida contratos heredados de configuraciones anteriores.
        # Los campos directos toman el nuevo default; los campos de asignación conservan
        # sus ajustes contextuales, ya normalizados al formato canónico.
        from apps.muestras.api.services.limit_contract_reconciliation import (
            reconcile_test_assignments,
            reconcile_test_predefined_details,
        )
        reconcile_test_assignments(prueba, reset_direct_defaults=True)
        reconcile_test_predefined_details(prueba, reset_direct_defaults=True)

        source.refresh_from_db()
        return Response(self.get_serializer(source).data, status=status.HTTP_200_OK)

    @action(detail=False, methods=["post"], url_path="resolve-preview")
    def resolve_preview(self, request):
        muestra = Muestra.objects.get(pk=request.data.get("muestra"))
        prueba = Prueba.objects.get(pk=request.data.get("prueba"))
        resultado = _get_optional(PruebaResultado, request.data.get("resultado"))
        division = _get_optional(PruebaResultadoDivision, request.data.get("division"))
        componente = _get_optional(PruebaResultadoComponente, request.data.get("componente"))
        valor = request.data.get("valor")
        if valor in [None, ""]:
            resolved = resolve_limit_for_result(muestra, prueba, resultado, division, componente)
        else:
            resolved = resolve_and_evaluate_result(valor, muestra, prueba, resultado, division, componente)
        return Response(resolved.to_dict())


class PruebaLimiteCampoViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "valores_limite.ver",
        "retrieve": "valores_limite.ver",
        "create": "valores_limite.editar_matriz",
        "update": "valores_limite.editar_matriz",
        "partial_update": "valores_limite.editar_matriz",
        "destroy": "valores_limite.editar_matriz",
        "restore": "valores_limite.editar_matriz",
    }
    queryset = PruebaLimiteCampo.objects.select_related(
        "fuente_limite", "fuente_limite__prueba", "fuente_limite__catalogo_fuente", "resultado", "division", "componente"
    )
    serializer_class = PruebaLimiteCampoSerializer
    search_fields = ["nombre", "codigo", "fuente_limite__prueba__nombre_variable"]
    ordering_fields = ["orden", "nombre", "codigo", "created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        fuente = self.request.query_params.get("fuente_limite")
        prueba = self.request.query_params.get("prueba")
        if fuente:
            queryset = queryset.filter(fuente_limite_id=fuente)
        if prueba:
            queryset = queryset.filter(fuente_limite__prueba_id=prueba)
        return queryset.order_by("orden", "id")


class CriterioEvaluacionLimiteViewSet(SoftDeleteViewSet):
    permission_action_map = {
        "list": "valores_limite.ver",
        "retrieve": "valores_limite.ver",
        "create": "valores_limite.editar_matriz",
        "update": "valores_limite.editar_matriz",
        "partial_update": "valores_limite.editar_matriz",
        "destroy": "valores_limite.editar_matriz",
        "restore": "valores_limite.editar_matriz",
    }
    queryset = CriterioEvaluacionLimite.objects.select_related(
        "fuente_limite",
        "fuente_limite__prueba",
        "fuente_limite__catalogo_fuente",
        "fuente_limite__campo_tecnico_muestra",
        "catalogo_item",
        "catalogo_item__catalogo",
        "escala_item",
        "escala_item__escala",
    )
    serializer_class = CriterioEvaluacionLimiteSerializer
    search_fields = ["nombre", "codigo", "fuente_limite__prueba__nombre_variable"]
    ordering_fields = ["orden", "nombre", "created_at"]

    def get_queryset(self):
        queryset = super().get_queryset()
        fuente = self.request.query_params.get("fuente_limite")
        prueba = self.request.query_params.get("prueba")
        if fuente:
            queryset = queryset.filter(fuente_limite_id=fuente)
        if prueba:
            queryset = queryset.filter(fuente_limite__prueba_id=prueba)
        return queryset


def _build_limit_fields_from_prueba(
    fuente,
    operador="max",
    unidad="",
    tipo_comparacion="numerica",
    escala_comparacion_id=None,
    modo_ordinal="orden",
    evaluacion_opciones=None,
):
    """
    Genera columnas de límite desde la estructura REAL de la prueba.

    Reglas:
    - Si la prueba tiene resultado simple: crea 1 campo.
    - Si la prueba tiene divisiones/componentes: crea 1 campo por hoja final.
      Ej: Espuma -> Secuencia I/FI, Secuencia I/EI, Secuencia II/FII...
    - Si el operador base es "between", no crea un campo ambiguo; crea dos campos
      por hoja final: mínimo y máximo. Así se puede corregir una asociación que
      antes se creó como solo máximo.
    - Al regenerar, desactiva los campos viejos de esa fuente que ya no aplican
      para evitar columnas basura como "viscosidad_v" cuando realmente eran
      Vmin/Vmax.
    """
    operador = _normalize_operator(operador)
    unidad = unidad or ""
    prueba = fuente.prueba

    leaves = _test_limit_leaves(prueba)
    wanted_codes = set()
    created_or_updated = []
    order = 1

    for leaf in leaves:
        if operador == "between":
            for suffix, op, label in [("min", "min", "mínimo"), ("max", "max", "máximo")]:
                nombre = f"{leaf['nombre']} {label}".strip()
                codigo = f"{leaf['codigo']}_{suffix}"
                field = _upsert_limit_field(
                    fuente=fuente,
                    nombre=nombre,
                    codigo=codigo,
                    operador=op,
                    unidad=unidad,
                    orden=order,
                    tipo_comparacion=tipo_comparacion,
                    escala_comparacion_id=escala_comparacion_id,
                    modo_ordinal=modo_ordinal,
                    evaluacion_opciones=evaluacion_opciones,
                    resultado=leaf.get("resultado"),
                    division=leaf.get("division"),
                    componente=leaf.get("componente"),
                )
                wanted_codes.add(codigo)
                created_or_updated.append(field)
                order += 1
        else:
            nombre = leaf["nombre"]
            codigo = leaf["codigo"]
            field = _upsert_limit_field(
                fuente=fuente,
                nombre=nombre,
                codigo=codigo,
                operador=operador,
                unidad=unidad,
                orden=order,
                tipo_comparacion=tipo_comparacion,
                escala_comparacion_id=escala_comparacion_id,
                modo_ordinal=modo_ordinal,
                evaluacion_opciones=evaluacion_opciones,
                resultado=leaf.get("resultado"),
                division=leaf.get("division"),
                componente=leaf.get("componente"),
            )
            wanted_codes.add(codigo)
            created_or_updated.append(field)
            order += 1

    # Importante para correcciones: si antes se generó un campo con el operador equivocado,
    # no lo dejamos activo ni visible en la matriz de valores.
    PruebaLimiteCampo.objects.filter(fuente_limite=fuente).exclude(codigo__in=wanted_codes).update(activo=False, deleted_at=timezone.now())

    return created_or_updated


def _normalize_operator(operador):
    return {
        "<=" : "max",
        ">=" : "min",
        "="  : "eq",
        "range": "between",
        "between": "between",
        "menor_igual": "max",
        "mayor_igual": "min",
        "igual": "eq",
        "entre": "between",
        "rango": "between",
        "min_max": "between",
    }.get(operador or "max", operador or "max")


def _upsert_limit_field(
    fuente,
    nombre,
    codigo,
    operador,
    unidad,
    orden,
    tipo_comparacion="numerica",
    escala_comparacion_id=None,
    modo_ordinal="orden",
    evaluacion_opciones=None,
    valor_global=None,
    resultado=None,
    division=None,
    componente=None,
):
    field, _ = PruebaLimiteCampo.objects.update_or_create(
        fuente_limite=fuente,
        codigo=codigo,
        defaults={
            "nombre": nombre or f"Campo {orden}",
            "operador": operador,
            "tipo_comparacion": tipo_comparacion or "numerica",
            "escala_comparacion_id": escala_comparacion_id if tipo_comparacion in ["escala", "escala_ordinal"] else None,
            "modo_ordinal": modo_ordinal or "orden",
            "evaluacion_opciones": evaluacion_opciones,
            "valor_global": valor_global if valor_global not in [None, ""] else None,
            "unidad": unidad or None,
            "resultado": resultado,
            "division": division,
            "componente": componente,
            "orden": orden,
            "activo": True,
            "deleted_at": None,
        },
    )
    return field


def _get_optional(model, pk):
    if not pk:
        return None
    return model.objects.get(pk=pk)
def _validate_semaphore_rule(rule, label, field_config=None):
    if not isinstance(rule, dict):
        return None

    field_config = field_config or {}
    if field_config.get("origen_limite") in {"asignacion", "informativo"}:
        return None
    operator = field_config.get("operador") or "max"
    comparison_type = field_config.get("tipo_comparacion") or "numerica"
    yellow = bool(rule.get("usar_amarillo"))
    required = {
        "min": ["min_critico", "min_aceptable"] if yellow else ["min_aceptable"],
        "between": (
            ["min_critico", "min_aceptable", "max_aceptable", "max_critico"]
            if yellow else ["min_aceptable", "max_aceptable"]
        ),
        "max": ["max_aceptable", "max_critico"] if yellow else ["max_aceptable"],
        "eq": ["valor_esperado"],
        "neq": ["valor_esperado"],
    }.get(operator, [])
    if any(rule.get(key) in [None, ""] for key in required):
        raise ValueError(f"{label}: complete todos los valores requeridos por el operador.")

    if comparison_type in {"escala", "escala_ordinal", "booleano"}:
        return None

    def number(key):
        value = rule.get(key)
        if value in [None, ""]:
            return None
        try:
            return Decimal(str(value))
        except (InvalidOperation, TypeError, ValueError):
            raise ValueError(f"{label}: {key} debe ser numerico.")

    minimum_critical = number("min_critico")
    minimum_ok = number("min_aceptable")
    maximum_ok = number("max_aceptable")
    maximum_critical = number("max_critico")
    if minimum_ok is not None and maximum_ok is not None and minimum_ok > maximum_ok:
        raise ValueError(f"{label}: el minimo aceptable no puede superar el maximo aceptable.")
    if rule.get("usar_amarillo") and minimum_critical is not None and minimum_ok is not None and minimum_critical > minimum_ok:
        raise ValueError(f"{label}: el minimo critico no puede superar el minimo aceptable.")
    if rule.get("usar_amarillo") and maximum_ok is not None and maximum_critical is not None and maximum_ok > maximum_critical:
        raise ValueError(f"{label}: el maximo aceptable no puede superar el maximo critico.")
    return None
